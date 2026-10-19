# Builds the style templates of the generator from the user's design files (run once when a design changes).
#   python3 tools/make_style_templates.py KP_2_v3.pptx KP_BMF720_D5_light_v3.pptx KP_BMF720_D5_light.pptx KP_NPL10720_D6_heavy_label.pptx
# The two light designs share KP_2_v3's layout (same shape ids) but were made from an older text version, so their
# texts are synced from KP_2_v3 (the approved wording); colours, pictures and layout stay the light ones.
import os, shutil, sys
from pptx import Presentation

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.normpath(os.path.join(HERE, '..', 'templates'))


def sync_texts(src, dst):
    a, b = Presentation(src), Presentation(dst)
    n = 0
    for sa, sb in zip(a.slides, b.slides):
        B = {sh.shape_id: sh for sh in sb.shapes}
        for sh in sa.shapes:
            o = B.get(sh.shape_id)
            if o is None or not sh.has_text_frame or not o.has_text_frame: continue
            if sh.text_frame.text == o.text_frame.text: continue
            pa, pb = sh.text_frame.paragraphs, o.text_frame.paragraphs
            if len(pa) != len(pb) or any(len(x.runs) != len(y.runs) for x, y in zip(pa, pb)): continue
            for x, y in zip(pa, pb):
                for ra, rb in zip(x.runs, y.runs):
                    if ra.text != rb.text: rb.text = ra.text; n += 1
    return b, n


def main(kp2, light_v3, light, d6):
    os.makedirs(OUT, exist_ok=True)
    for src, name in ((light_v3, 'KP_light_v3.pptx'), (light, 'KP_light.pptx')):
        prs, n = sync_texts(kp2, src)
        prs.save(os.path.join(OUT, name)); print(name, 'texts synced:', n)
    shutil.copyfile(d6, os.path.join(OUT, 'KP_npl_label.pptx')); print('KP_npl_label.pptx')


if __name__ == '__main__':
    main(*sys.argv[1:5])
