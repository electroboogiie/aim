# Builds templates/KP_crane_label.pptx — the «Heavy Label» design for a crane in a kit — from the user's НПЛ-10 design
# (templates/KP_npl_label.pptx). Run once when either design changes:   python3 tools/make_crane_label.py
# The НПЛ-10 pages (dimensions, НПЛ-10 specs) are dropped and pages renumbered 01–08; page 02 becomes the kit page;
# fixed wording and headline pictures are rewritten for a BMF 720 kit. The engine (kpgen/d6.py, style crane_label)
# then writes the kit's own numbers, model headlines, crane cut-out and photos.
import os, re, sys
from pptx import Presentation

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.normpath(os.path.join(HERE, '..')))
from kpgen import d6  # noqa: E402
from kpgen.d6 import INK, OR, WHITE, OUTLINE, set_headline, set_cutout, hero_photo  # noqa: E402
from kpgen.cache import Cache  # noqa: E402
from kpgen.pptx_util import EMU, by_id, by_name, put, set_lines, set_image, delete_slide, prune_images  # noqa: E402

TPL = os.path.normpath(os.path.join(HERE, '..', 'templates'))
KEEP = [0, 1, 4, 5, 6, 7, 8, 9]   # slides of the НПЛ-10 design kept, in order


def text(s, sid, *paras):
    put(by_id(s, sid), *paras)


def main():
    prs = Presentation(os.path.join(TPL, 'KP_npl_label.pptx'))
    S = list(prs.slides)
    for i in (3, 2): delete_slide(prs, S[i])
    S = list(prs.slides); n = len(S)
    cache = Cache(None)
    for no, s in enumerate(S, 1):
        for sh in s.shapes:
            if not sh.has_text_frame: continue
            t = sh.text_frame.text
            if t == 'НПЛ-10 × BMF 720': put(sh, 'BMF 720 · КРАН У КОМПЛЕКТІ')
            elif re.fullmatch(r'\d\d / 10', t): put(sh, f'{no:02d} / {n:02d}')
            elif re.fullmatch(r'\d\d', t) and sh.top / EMU > 11: put(sh, f'{no:02d}')
            elif re.match(r'\d\d // ', t):
                r = sh.text_frame.paragraphs[0].runs[0]; r.text = f'{no - 1:02d}' + r.text[2:]
    cover, kit, crane, spec, price, terms, adv, end = S
    # 01 cover
    text(cover, 11, 'КРАН-МАНІПУЛЯТОР'); text(cover, 12, 'У КОМПЛЕКТІ')
    text(cover, 14, 'З РОТАТОРОМ 4 т І ГРЕЙФЕРОМ 0,24 м²')
    text(cover, 17, ['60', '  кНм']); text(cover, 18, 'ПІДЙОМНИЙ МОМЕНТ')
    text(cover, 23, ['1100', '  кг']); text(cover, 24, 'ВАНТАЖ НА ВИЛЬОТІ 4 м')
    text(cover, 26, ['370', '  °']); text(cover, 27, 'ПОВОРОТ СТРІЛИ')
    set_headline(cache, by_name(cover, 'hdr_cover_title')[0], [[('BMF ', INK), ('720', OR)]], [[('НПЛ-', INK), ('10', OR)]],
                 maxw=int(7.34 * EMU))
    set_headline(cache, by_name(cover, 'hdr_cover_10t')[0], [[('720', OUTLINE)]], [[('10т', OUTLINE)]], outline=4)
    set_headline(cache, by_name(cover, 'hdr_cover_slogan')[0], [[('7,2 м ВИЛЬОТУ.', WHITE)], [('ЖОДНИХ КОМПРОМІСІВ.', OR)]],
                 [[('10 ТОНН ЛІСУ.', WHITE)], [('ЖОДНИХ КОМПРОМІСІВ.', OR)]])
    set_cutout(cache, by_name(cover, 'hero_cutout')[0], '720', (-0.1, 4.75, 5.75, 4.3))
    # 02 kit
    set_headline(cache, by_name(kit, 'hdr_p2_hdr')[0], [[('ГОТОВИЙ', WHITE)], [('ДО МОНТАЖУ.', OR)]],
                 [[('ЗБУДОВАНИЙ', WHITE)], [('ДЛЯ БЕЗДОРІЖЖЯ.', OR)]])
    text(kit, 25, '01 // СКЛАД КОМПЛЕКТУ · BMF 720')
    text(kit, 28, ['2', '  поз.']); text(kit, 29, 'У КОМПЛЕКТІ')
    text(kit, 31, ['4', '  т']); text(kit, 32, 'РОТАТОР')
    text(kit, 34, ['0,24', '  м²']); text(kit, 35, 'ГРЕЙФЕР')
    text(kit, 37, ['675', '  кг']); text(kit, 38, 'ВАГА КРАНА')
    text(kit, 40, 'BMF 720 у комплекті: Кран-маніпулятор BMF 720, платформа з аутригерами. Ротатор 4 т і грейфер 0,24 м² '
                  'уже входять у ціну крана.')
    text(kit, 41, 'СКЛАД КОМПЛЕКТУ')
    set_lines(by_id(kit, 42), ['01  Кран-маніпулятор BMF 720 — у комплекті з ротатором 4 т і захватом 0,24 м²',
                               '02  Платформа з аутригерами — опорні лапи, РВТ, з’єднувальні болти, фланці'])
    text(kit, 43, 'ДОДАТКОВО ЗА ЗАПИТОМ')
    set_lines(by_id(kit, 44), ['Регулюючий клапан — ціна за запитом', 'Грейфери Hardox — ціна за запитом',
                               'Ротатори — ціна за запитом', 'Опора F/D-типу («метелик») — ціна за запитом'])
    text(kit, 52, 'BMF 905')
    text(kit, 54, 'ЧОМУ BMF')
    text(kit, 55, 'Ми — офіційні дистриб’ютори BMF в Україні. Підберемо кран, ротатор, грейфер і основу під ваш причіп '
                  'або трактор, а всі умови узгодимо до підписання договору.')
    pic = by_id(kit, 3)
    set_image(pic, hero_photo('bmf/720_workshop.jpg', pic.width / EMU, pic.height / EMU, 0.5, 0.62))
    # 05 price
    text(price, 53, 'BMF 720 У КОМПЛЕКТІ')
    pic = by_id(price, 3)
    set_image(pic, hero_photo('prod/base_outriggers.jpg', pic.width / EMU, pic.height / EMU, 0.5, 0.55))
    # 06 terms
    text(terms, 49, 'поставки під замовлення з моменту підписання договору')
    pic = by_id(terms, 50)
    set_image(pic, d6.photo_part('prod/tank.jpg', pic.width / EMU, pic.height / EMU))
    text(terms, 59, 'ГІДРАВЛІЧНИЙ БАК В ЗБОРІ')
    # 07 advantages
    set_headline(cache, by_name(adv, 'hdr_p9_hdr')[0], [[('ЧОМУ НАС', WHITE)], [('ОБИРАЮТЬ.', OR)]],
                 [[('ЧОМУ НПЛ-10', WHITE)], [('ОБИРАЮТЬ.', OR)]])
    # 08 end
    set_headline(cache, by_name(end, 'hdr_p10_hdr')[0], [[('ЗАМОВТЕ', INK)], [('СВІЙ BMF 720.', WHITE)]],
                 [[('ЗАМОВТЕ', INK)], [('СВІЙ НПЛ-10.', WHITE)]])
    set_cutout(cache, by_name(end, 'cta_cutout')[0], '720', (4.0, 5.3, 3.86, 3.4))
    prune_images(prs)
    out = os.path.join(TPL, 'KP_crane_label.pptx')
    prs.save(out); print(out)


if __name__ == '__main__':
    main()
