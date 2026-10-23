# PDF export through LibreOffice. Each run gets its own LibreOffice profile so parallel runs do not
# block each other on the shared user profile.
import os, pathlib, shutil, signal, subprocess, tempfile

WIN_PATHS = [os.path.join(os.environ.get(v, ''), 'LibreOffice', 'program', 'soffice.exe')
             for v in ('ProgramFiles', 'ProgramFiles(x86)', 'ProgramW6432') if os.environ.get(v)]


class PdfError(Exception):
    pass


def soffice_bin():
    for name in (os.environ.get('KPGEN_SOFFICE'), 'soffice', 'libreoffice'):
        if name and shutil.which(name): return shutil.which(name)
    for path in WIN_PATHS:
        if os.path.isfile(path): return path
    raise PdfError('Не знайдено LibreOffice (soffice) для створення PDF.')


def to_pdf(pptx, out_dir, timeout=None):
    return convert(pptx, out_dir, 'pdf', timeout)


def to_png(pptx, out_dir, timeout=None):
    """PNG of the first slide (LibreOffice exports only the first page to PNG): the KP cover preview."""
    return convert(pptx, out_dir, 'png', timeout)


def convert(pptx, out_dir, fmt, timeout=None):
    timeout = timeout or int(os.environ.get('KPGEN_PDF_TIMEOUT', '180'))
    profile = tempfile.mkdtemp(prefix='kpgen_lo_')
    try:
        pdf = os.path.join(out_dir, os.path.splitext(os.path.basename(pptx))[0] + '.' + fmt)
        if os.path.exists(pdf): os.remove(pdf)   # soffice may exit 0 without writing: never report a stale file
        cmd = [soffice_bin(), '-env:UserInstallation=' + pathlib.Path(profile).as_uri(), '--headless', '--norestore',
               '--convert-to', fmt, '--outdir', out_dir, pptx]
        # own session: on timeout the whole LibreOffice tree (wrapper + soffice.bin) is killed
        p = subprocess.Popen(cmd, stdout=subprocess.PIPE, stderr=subprocess.PIPE, text=True,
                             encoding='utf-8', errors='replace', start_new_session=True)
        try:
            out, err = p.communicate(timeout=timeout)
        except subprocess.TimeoutExpired:
            try:
                if hasattr(os, 'killpg'): os.killpg(p.pid, signal.SIGKILL)
                else: p.kill()
            except ProcessLookupError: pass
            p.communicate()
            raise PdfError(f'LibreOffice не встиг створити {fmt.upper()} за {timeout} с.')
        if p.returncode != 0 or not os.path.exists(pdf):
            raise PdfError(f'LibreOffice не створив {fmt.upper()} (код {p.returncode}): {(err or out).strip()[-500:]}')
        return pdf
    finally:
        shutil.rmtree(profile, ignore_errors=True)


def _page_key(slide, width):
    """Hash of a page: its XML and the parts it uses (pictures, layout). Same hash = same picture."""
    import hashlib
    h = hashlib.sha1(str(width).encode())
    h.update(slide.part.blob)
    for rel in sorted(slide.part.rels.values(), key=lambda r: r.rId):
        if not rel.is_external: h.update(rel.target_part.blob)
    return h.hexdigest()[:20]


def pages_png(pptx, out_dir, width=640, timeout=None, cache_dir=None):
    """PNG of every page (LibreOffice exports only the first slide to PNG: each page goes into its own one-slide pptx,
    all converted in one LibreOffice run). With cache_dir, pages that did not change since an earlier run are not
    converted again. Returns the PNG paths in page order."""
    from pptx import Presentation
    src = Presentation(pptx); slides = list(src.slides); n = len(slides)
    keys = [_page_key(s, width) for s in slides] if cache_dir else [None] * n
    if cache_dir: os.makedirs(cache_dir, exist_ok=True)
    pngs = [os.path.join(out_dir, f'page{i + 1:02d}.png') for i in range(n)]
    todo = []
    for i in range(n):
        hit = cache_dir and os.path.join(cache_dir, keys[i] + '.png')
        if hit and os.path.exists(hit): shutil.copyfile(hit, pngs[i])
        else: todo.append(i)
    parts = []
    for i in todo:
        p = Presentation(pptx); ids = p.slides._sldIdLst
        for k, el in enumerate(list(ids)):
            if k != i:
                p.part.drop_rel(el.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id')); ids.remove(el)
        part = os.path.join(out_dir, f'page{i + 1:02d}.pptx'); p.save(part); parts.append(part)
    if not parts: return pngs
    h = int(width * 11.69 / 8.27)
    flt = 'png:impress_png_Export:{"PixelWidth":{"type":"long","value":%d},"PixelHeight":{"type":"long","value":%d}}' % (width, h)
    timeout = timeout or int(os.environ.get('KPGEN_PDF_TIMEOUT', '180'))
    profile = tempfile.mkdtemp(prefix='kpgen_lo_')
    try:
        cmd = [soffice_bin(), '-env:UserInstallation=' + pathlib.Path(profile).as_uri(), '--headless', '--norestore',
               '--convert-to', flt, '--outdir', out_dir, *parts]
        p = subprocess.Popen(cmd, stdout=subprocess.PIPE, stderr=subprocess.PIPE, text=True, encoding='utf-8', errors='replace',
                             start_new_session=True)
        try: p.communicate(timeout=timeout)
        except subprocess.TimeoutExpired:
            try:
                if hasattr(os, 'killpg'): os.killpg(p.pid, signal.SIGKILL)
                else: p.kill()
            except ProcessLookupError: pass
            p.communicate()
            raise PdfError(f'LibreOffice не встиг створити перегляд сторінок за {timeout} с.')
    finally:
        shutil.rmtree(profile, ignore_errors=True)
        for part in parts:
            try: os.remove(part)
            except OSError: pass
    if not all(os.path.exists(pngs[i]) for i in todo): raise PdfError('LibreOffice не створив перегляд сторінок.')
    if cache_dir:
        for i in todo: shutil.copyfile(pngs[i], os.path.join(cache_dir, keys[i] + '.png'))
        old = sorted((os.path.join(cache_dir, f) for f in os.listdir(cache_dir)), key=os.path.getmtime)[:-400]
        for f in old:
            try: os.remove(f)
            except OSError: pass
    return pngs
