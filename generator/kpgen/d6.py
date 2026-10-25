# «НПЛ-10 Heavy Label» style (the user's KP_NPL10720_D6_heavy_label.pptx, 10 pages): the design file is the template,
# texts and numbers of the kit are written into it by shape id. Only НПЛ-10 + crane kits.
# 21.10.2026, the user's remarks: no «10т» outline behind the cover figures, the crane's detailed manufacturer drawing
# (KP_2_v3) instead of the line drawing, the load chart with the crane, larger price lines, a larger НПЛ-10 on the
# last page; headlines with the crane's reach and torque are redrawn for each crane.
import io, os, re
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Pt

from . import raster as R
from .cache import Cache
from .pptx_util import EMU, I, by_id, by_name, put, set_lines, set_image, remove, prune_images
from .styles import template_path
from .profiles import PROFILES, NPL10
from .spec import SpecError, client_line

INK, OR = (13, 14, 16), (242, 83, 10)
FONT_BLACK = next((f for f in (os.path.join(R.ASSETS, 'fonts', 'InterDisplay-Black.otf'),
                               '/usr/share/fonts/opentype/inter/InterDisplay-Black.otf') if os.path.isfile(f)), R.FONT_B)
PAGES = ['cover', 'npl_info', 'npl_dims', 'npl_spec', 'crane', 'crane_spec', 'price', 'terms', 'advantages', 'end']


def display(text, pt=46, color=INK, dpi=300):
    f = ImageFont.truetype(FONT_BLACK, int(pt / 72 * dpi)); tr = -0.02 * f.size
    w = int(sum(f.getlength(c) for c in text) + tr * len(text) + f.size * 0.2); h = int(f.size * 1.3)
    im = Image.new('RGBA', (w, h), (0, 0, 0, 0)); d = ImageDraw.Draw(im); x = 0
    for c in text:
        d.text((x, 0), c, font=f, fill=color + (255,)); x += f.getlength(c) + tr
    return im.crop(im.getbbox())


def png(im):
    b = io.BytesIO(); im.save(b, 'PNG'); return b.getvalue()


class D6:
    def __init__(self, spec, build_dir):
        if spec['kit_type'] != 'npl_crane':
            raise SpecError('Стиль «НПЛ-10 Heavy Label» є лише для комплектів НПЛ-10 + кран. Оберіть інший стиль.')
        self.spec = spec; self.c = spec['crane']; self.cs = self.c['specs']
        self.cache = Cache(build_dir); self.build_dir = build_dir
        R.set_theme('dark')

    def build(self):
        sp, cs, m = self.spec, self.cs, self.c['model']
        prs = Presentation(template_path('npl_label')); S = list(prs.slides); self.pages = dict(zip(PAGES, S))
        g = lambda k, d='—': cs.get(k) or d
        old = 'BMF 720'
        # every page: header "НПЛ-10 × BMF 720" and other mentions of the template's crane
        for s in S:
            for sh in s.shapes:
                if sh.has_text_frame and old in sh.text_frame.text and m != old:
                    for p in sh.text_frame.paragraphs:
                        for r in p.runs:
                            if old in r.text: r.text = r.text.replace(old, m)
        ct = sp['contacts']; npl = sp.get('npl') or {}
        # 01 cover
        s = S[0]
        remove(by_id(s, 5))   # «10т» outline behind the figures
        put(by_id(s, 6), f"КОМЕРЦІЙНА ПРОПОЗИЦІЯ № {sp['number']}")
        cl = client_line(sp, 'ДЛЯ')
        put(by_id(s, 7), f"{sp['date']}  ·  {cl[0].upper()}: {cl[1].upper()}" if cl else f"ВЕРСІЯ {sp['date']}  ·  САРНИ, UA")
        put(by_id(s, 20), [cs['reach_m'], None]); put(by_id(s, 26), [npl.get('stakes') or '4', None])
        self.contacts(s, (31, 33, 35, 37))
        put(by_id(s, 47), sp['maker']['line'].upper())
        # 04 НПЛ-10 specs
        s = S[3]
        put(by_id(s, 32), f"{npl.get('stakes') or '4'} шт"); put(by_id(s, 50), f"{npl.get('wheels') or '500/45-22,5'} · 4 шт")
        put(by_name(s, 'fld_npl_11')[0], npl.get('brakes') or 'гідравлічна')
        lab = [sh for sh in s.shapes if sh.has_text_frame and sh.text_frame.text.startswith('Колесо ходове з диском')]
        if lab: put(lab[0], f"Колесо ходове з диском {npl.get('wheels') or '500/45-22,5'}")
        # 05 crane
        s = S[4]
        prof = PROFILES.get(self.c.get('profile')) or {}
        native = m == old   # the template's texts are about BMF 720: other cranes take their profile texts
        lead, lead2 = self.c.get('lead') or (None if native else prof.get('lead')), self.c.get('lead2') or (None if native else prof.get('lead2'))
        if lead: put(by_id(s, 28), lead)
        if lead2: put(by_id(s, 29), lead2)
        put(by_id(s, 49), cs['reach_m']); put(by_id(s, 51), f"{cs['turn_deg']}°")
        # 06 crane specs
        s = S[5]
        put(by_id(s, 27), [cs['reach_m'], None]); put(by_id(s, 30), [cs['torque_knm'], None])
        put(by_id(s, 33), [cs['turn_deg'], None]); put(by_id(s, 36), [g('pressure_bar'), None])
        put(by_id(s, 66), f"РОТАТОР {cs['rotator_t']} т   +   ГРЕЙФЕР {cs['grapple_m2']} м²")
        flow = g('valve_flow_l'); flow = flow if flow.endswith('л') or flow == '—' else flow + 'l'
        vals = [cs['reach_m'], cs['torque_knm'], cs['lift_4m_kg'], cs['lift_full_kg'], cs['grapple_m2'], cs['rotator_t'] + 'т',
                f"{g('valve')} ({flow})", g('slew_knm'), g('slew_cyl'), cs['turn_deg'], g('pressure_bar'), g('pump_lpm'),
                '+' if g('wide_angle', '+') in ('так', '+') else '—', '+' if g('telescope', '+') in ('так', '+') else '—', g('weight_kg')]
        for i, v in enumerate(vals):
            f = by_name(s, f'fld_bmf_{i}')
            if f: put(f[0], v)
        # 07 price
        self.price(S[6])
        # 08 terms
        s = S[7]
        for (tid, xid), t in zip(((29, 30), (33, 34), (37, 38), (41, 42), (45, 46)), sp['terms']):
            put(by_id(s, tid), t['title'].upper()); put(by_id(s, xid), t['text'])
        # 10 contacts
        s = S[9]
        self.contacts(s, (34, 37, 40, 43))
        for i in (34, 37, 40, 43): by_id(s, i).width = I(2.75)
        self.npl_photo(by_id(s, 31), 3.45, 4.55, 4.6, 4.65)
        self.npl_photo(by_name(S[0], 'hero_cutout')[0], -0.05, 3.85, 5.29, 5.35)
        put(by_id(s, 57), sp['maker']['line'].upper())
        self.crane_art(S)
        self.overrides()
        prune_images(prs)
        return prs

    def crane_art(self, S):
        """Page 05: the manufacturer side drawing of the crane; page 06: load chart with the crane; headlines with
        the crane's reach and torque."""
        from .ultra import v3_pieces, drawing_panel, load_chart, rtext
        cs = self.cs
        e, parts, dims, side = v3_pieces(self.spec, self.build_dir)
        R.set_theme('dark')
        # 05: side view with its dimensions in the «схема вильоту» frame
        p0 = parts[0]; dd = [d for d in dims if d[2] < p0[2] + p0[4] - 0.15]
        data = self.cache.get('d6-crane-dwg', drawing_panel, ([p0], dd))
        s = S[4]; old = by_name(s, 'dwg_crane')[0]
        im = Image.open(io.BytesIO(data)); k = min(5.35 / im.width, 2.1 / im.height)
        pic = s.shapes.add_picture(io.BytesIO(data), I(0.7), I(8.32 + (2.1 - im.height * k) / 2), I(im.width * k), I(im.height * k))
        old._element.addprevious(pic._element); remove(old)
        # 06: chart with the crane on a dark card in place of the white bar chart
        s = S[5]
        reach = float(cs['reach_m'].replace(',', '.'))
        data = self.cache.get('d6-chart', load_chart, (side, e.span, reach, int(cs['lift_4m_kg']), int(cs['lift_full_kg'])))
        box = by_id(s, 38); box.fill.solid(); box.fill.fore_color.rgb = RGBColor(*INK)
        for sid in range(40, 63): remove(by_id(s, sid))
        title = by_id(s, 39)
        for p in title.text_frame.paragraphs:
            for r_ in p.runs: r_.font.color.rgb = RGBColor(156, 162, 172)
        im = Image.open(io.BytesIO(data)); k = min(7.0 / im.width, 1.88 / im.height); w, h = im.width * k, im.height * k
        pic = s.shapes.add_picture(io.BytesIO(data), I(0.5 + (7.26 - w) / 2), I(3.3 + 2.25 - h - 0.06), I(w), I(h))
        box._element.addnext(pic._element)
        # headlines drawn with the template's numbers (BMF 720): redraw for another reach / torque
        for slide, name, lines, cap, pitch, top_px, want in (
                (S[4], 'hdr_p5_hdr', [[('СИЛА BMF.', (255, 255, 255))], [(f"ВИЛІТ {cs['reach_m']} м.", OR)]], 118, 1.24, 2, cs['reach_m'] != '7,2'),
                (S[5], 'hdr_p6_hdr', [[('ПІДЙОМНИЙ МОМЕНТ', INK)], [(f"{cs['torque_knm']} кНм.", OR)]], 99, 1.06, 48, cs['torque_knm'] != '60')):
            if not want: continue
            sh = by_name(slide, name)[0]
            res = self.cache.get('d6-hdr', lambda L, c, pi: (lambda r: (png(r[0]), [r[2]]))(rtext(L, FONT_BLACK, c, pi, -0.02)), (lines, cap, pitch))
            data, (off,) = res; im = Image.open(io.BytesIO(data))
            top = sh.top / EMU + (top_px - off) / 300
            set_image(sh, data); sh.top, sh.width, sh.height = I(top), I(im.width / 300), I(im.height / 300)

    def contacts(self, s, ids):
        ct = self.spec['contacts']
        ph, site, mail, addr = (by_id(s, i) for i in ids)
        if ct['phones']: set_lines(ph, ct['phones'][:2])
        if ct['sites']: set_lines(site, ct['sites'][:2])
        if ct['email']: set_lines(mail, [ct['email']])
        if ct['address']: set_lines(addr, ct['address'].split('\n'))

    def price(self, s):
        sp = self.spec; rows = sp['rows']
        slots = [(31, 32, 33, 34), (35, 36, 37, 38), (39, 40, 41, 42), (43, 44, 45, 46)]
        shown = rows[:4]
        if len(rows) > 4:   # the receipt has four lines: the rest go into the last one
            rest = rows[3:]
            shown = rows[:3] + [dict(no=4, title='Також: ' + ', '.join(r['title'] for r in rest), desc='', price_text='')]
        n = len(shown); pitch = {1: 0.95, 2: 0.95, 3: 0.78}.get(n, 0.59)
        pts = {1: (11, 16, 10.5, 11), 2: (11, 16, 10.5, 11), 3: (10, 14, 9.5, 10)}.get(n, (8.5, 12.5, 8.5, 9))
        for k, (no, title, desc, price) in enumerate(slots):
            if k >= len(shown):
                for i in (no, title, desc, price): remove(by_id(s, i))
                continue
            y = 5.35 + k * pitch
            for sid, pt, (x, dy, w, h) in zip((no, title, desc, price), pts,
                                              ((1.1, 0.02, 0.4, 0.3), (1.55, 0, 4.2, 0.35), (1.55, pts[1] / 72 * 1.45, 4.2, 0.4), (5.3, 0.03, 1.86, 0.3))):
                sh = by_id(s, sid); sh.left, sh.top, sh.width, sh.height = I(x), I(y + dy), I(w), I(h)
                for p in sh.text_frame.paragraphs:
                    for r_ in p.runs: r_.font.size = Pt(pt)
            r = shown[k]
            put(by_id(s, no), f'{k + 1:02d}'); put(by_id(s, title), r['title']); put(by_id(s, desc), r['desc'] or ' ')
            pt = r['price_text'] or ''
            put(by_id(s, price), '✓ В КОМПЛЕКТІ' if not pt or 'комплект' in pt.lower() else pt)
        put(by_id(s, 29), f"КП № {sp['number']} · {sp['date']}")
        total = re.sub(r'\s*₴\s*$', '', sp['total_text']).strip()
        cur = by_id(s, 50)
        if not re.search(r'\d', total):   # "ціна за запитом"
            cur.text_frame.paragraphs[0].runs[0].text = ''
        pic = by_name(s, 'fld_price')[0]
        im = self.cache.get('d6-price', lambda t: png(display(t)), (total,))
        img = Image.open(io.BytesIO(im)); h = pic.height; w = int(h * img.width / img.height)
        set_image(pic, im); pic.width = w
        cur.left = pic.left + w + int(0.1 * 914400)
        note = [f"Ціна вказана із урахуванням ПДВ. Станом на {sp.get('price_date') or sp['date']}."]
        if sp.get('valid_until'): note.append(f"Пропозиція дійсна до {sp['valid_until']}.")
        if sp.get('price_note'): note.append(sp['price_note'])
        set_lines(by_name(s, 'fld_price_note')[0], [' '.join(note)])
        put(by_id(s, 53), sp['title'].upper())

    def npl_photo(self, pic, x, y, w, h):
        """The НПЛ-10 cut-out (25.10.2026: the user's render) fitted into the template's box, standing on its bottom edge."""
        path = NPL10['cutout']
        data = self.cache.get('d6-npl', lambda p_: png(R.cutout_glow(R.alpha_crop(Image.open(R.asset(p_))), a=55)), (path,), (path,))
        im = Image.open(io.BytesIO(data)); k = min(w / im.width, h / im.height); pw, ph = im.width * k, im.height * k
        set_image(pic, data); pic.left, pic.top, pic.width, pic.height = I(x + (w - pw) / 2), I(y + h - ph), I(pw), I(ph)

    def overrides(self):
        ov = self.spec.get('overrides') or {}
        for key, o in (ov.get('bg') or {}).items():
            s = self.pages.get(key); path = (o or {}).get('path')
            if s is None or not path or not R.asset_ok(path): continue
            bgs = by_name(s, 'bg_grid')
            if not bgs: continue
            dark = sum(Image.open(io.BytesIO(bgs[0].image.blob)).convert('L').resize((8, 8)).getdata()) / 64 < 128
            R.set_theme('dark' if dark else 'light')
            data = self.cache.get(f"d6-underlay-{'dark' if dark else 'light'}", R.underlay_on,
                                  (bgs[0].image.blob, path, float(o.get('opacity', 0.18)), float(o.get('fx', 0.5)), float(o.get('fy', 0.5))), (path,))
            set_image(bgs[0], data)
        R.set_theme('dark')
        for key, texts in (ov.get('texts') or {}).items():
            s = self.pages.get(key)
            if s is None: continue
            for sid, text in (texts or {}).items():
                try: sh = by_id(s, int(sid))
                except (KeyError, ValueError): continue
                if sh.has_text_frame and text is not None: set_lines(sh, str(text).split('\n'))
