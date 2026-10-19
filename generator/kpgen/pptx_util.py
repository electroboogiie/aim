# python-pptx helpers for editing the KP_2_v3 template in place.
import copy, io
from pptx.oxml.ns import qn

EMU = 914400


def I(v):
    """Inches -> EMU."""
    return int(round(v * EMU))


def by_id(slide, i):
    for sh in slide.shapes:
        if sh.shape_id == i: return sh
    raise KeyError(f'shape {i}')


def by_name(slide, n): return [sh for sh in slide.shapes if sh.name == n]
def one(slide, n): return by_name(slide, n)[0]
def remove(sh): sh._element.getparent().remove(sh._element)


def put(sh, *paras):
    """Replace text paragraph by paragraph; each paragraph keeps the formatting of its first run.
    A paragraph given as a list sets its runs one by one (value + unit runs). None keeps it.
    A paragraph that already reads the new text is left alone (keeps the template's runs)."""
    ps = sh.text_frame.paragraphs
    for p, new in zip(ps, paras):
        if new is None: continue
        runs = p.runs
        if isinstance(new, (list, tuple)):
            for r, t in zip(runs, new):
                if t is not None and r.text != t: r.text = t
            continue
        if ''.join(r.text for r in runs) == new: continue
        runs[0].text = new
        for r in runs[1:]: r._r.getparent().remove(r._r)


def set_lines(sh, lines):
    """One paragraph per line: extra template paragraphs are dropped, missing ones cloned from the last."""
    txb = sh.text_frame._txBody
    ps = list(sh.text_frame.paragraphs)
    lines = list(lines) or ['']
    while len(ps) < len(lines):
        el = copy.deepcopy(ps[-1]._p); ps[-1]._p.addnext(el)
        ps = list(sh.text_frame.paragraphs)
    for p in ps[len(lines):]: txb.remove(p._p)
    put(sh, *lines)


def set_image(pic, data):
    """Point a picture at a new image part (never rewrites a part another slide may share)."""
    _, rId = pic.part.get_or_add_image_part(io.BytesIO(data))
    pic._element.blipFill.blip.set(qn('r:embed'), rId)


def fit(pic, im, box=None):
    """Fit the picture into its current box (or box=(x,y,w,h) in EMU), aspect kept, centred."""
    x, y, w, h = box or (pic.left, pic.top, pic.width, pic.height)
    ar = im.height / im.width
    nw, nh = (w, int(w * ar)) if w * ar <= h else (int(h / ar), h)
    pic.left, pic.top, pic.width, pic.height = x + (w - nw) // 2, y + (h - nh) // 2, nw, nh


def geom(sh, x, y, w, h): sh.left, sh.top, sh.width, sh.height = I(x), I(y), I(w), I(h)


def next_id(slide): return max(x.shape_id for x in slide.shapes) + 1


def clone(slide, proto):
    """Deep copy of a shape appended on top of the slide, with a fresh shape id."""
    el = copy.deepcopy(proto._element); slide.shapes._spTree.append(el)
    el.xpath('./*[1]/p:cNvPr')[0].set('id', str(next_id(slide)))
    return list(slide.shapes)[-1]


def dup_slide(prs, src_slide, index):
    """Copy of a slide (shapes and image relations) inserted at index."""
    new = prs.slides.add_slide(src_slide.slide_layout)
    tree = new.shapes._spTree
    for el in list(tree):
        if el.tag.split('}')[1] in ('sp', 'pic', 'grpSp', 'cxnSp', 'graphicFrame'): tree.remove(el)
    rid = {}
    for r in src_slide.part.rels.values():
        if r.reltype.endswith(('/slideLayout', '/notesSlide')): continue
        rid[r.rId] = new.part.relate_to(r._target, r.reltype)
    for el in src_slide.shapes._spTree:
        if el.tag.split('}')[1] in ('sp', 'pic', 'grpSp', 'cxnSp', 'graphicFrame'):
            c = copy.deepcopy(el)
            for e in c.iter():
                for k in (qn('r:embed'), qn('r:link'), qn('r:id')):
                    if e.get(k) in rid: e.set(k, rid[e.get(k)])
            tree.append(c)
    ids = prs.slides._sldIdLst; el = list(ids)[-1]; ids.remove(el); ids.insert(index, el)
    return new


def delete_slide(prs, slide):
    ids = prs.slides._sldIdLst
    for el in list(ids):
        if prs.slides.part.related_part(el.rId) is slide.part:
            prs.part.drop_rel(el.rId); ids.remove(el); return
    raise KeyError('slide')


def dim_box(slide, proto, x, y, t):
    """Native dimension label cloned from the user's dim_* text box style, centred at (x, y) inches."""
    el = copy.deepcopy(proto._element); slide.shapes._spTree.append(el)
    el.nvSpPr.cNvPr.set('id', str(next_id(slide)))
    sh = list(slide.shapes)[-1]
    sh.left, sh.top = I(x - sh.width / EMU / 2), I(y - sh.height / EMU / 2)
    put(sh, t); sh.name = 'dim_' + t
    return sh


def set_size(sh, hundredths):
    """Font size (in 1/100 pt) of every run in the shape."""
    for p in sh.text_frame.paragraphs:
        for r in p.runs: r._r.get_or_add_rPr().set('sz', str(hundredths))


def prune_images(prs):
    """Drop image relations no shape uses any more (pictures replaced by set_image keep the old one,
    which roughly doubles the file). Returns the number of dropped relations."""
    n = 0
    for s in prs.slides:
        xml = s._element.xml
        for rId, rel in list(s.part.rels.items()):
            if rel.reltype.endswith('/image') and f'"{rId}"' not in xml:
                s.part.rels.pop(rId); n += 1
    return n
