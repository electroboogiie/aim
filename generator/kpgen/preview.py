# Admin previews: how a photo will look in the KP after the style's treatment (page backgrounds, position tiles,
# option cards). Same raster functions as the generator, smaller output. Not thread-safe (raster THEME is global).
import io
from PIL import Image
from pptx import Presentation

from . import raster as R
from .styles import HERO_PAGES, STYLES, pages, template_path

_BG = {}


def _template_bg(style, key):
    """Image bytes of the page's background picture in the style's template (None when the page has none)."""
    if (style, key) not in _BG:
        origin = {k: i for k, _, i in pages(style)}
        s = list(Presentation(template_path(style)).slides)[origin[key]]
        _BG[(style, key)] = next((sh.image.blob for sh in s.shapes if sh.name in ('bg', 'bg_grid')), None)
    return _BG[(style, key)]


def _theme(style, base):
    if STYLES[style]['engine'] == 'd6' and base:   # D6 pages are dark or light: judge by the background itself
        dark = sum(Image.open(io.BytesIO(base)).convert('L').resize((8, 8)).getdata()) / 64 < 128
        return 'dark' if dark else 'light'
    return STYLES[style]['theme']


def _thumb(data, width):
    im = Image.open(io.BytesIO(data)).convert('RGB'); im.thumbnail((width, width * 2), Image.LANCZOS)
    b = io.BytesIO(); im.save(b, 'JPEG', quality=84); return b.getvalue()


def page(style, key, path, fx=0.5, fy=0.45, opacity=0.18, width=480):
    """JPEG of the page background with the admin photo, treated exactly as the generator does."""
    S = STYLES[style]
    base = _template_bg(style, key)
    R.set_theme(_theme(style, base))
    try:
        hero = S['engine'] == 'v3' and key in HERO_PAGES and not (key == 'cover' and S.get('cover') == 'panel')
        if hero:
            fn = R.cover_bg if key == 'cover' else R.end_bg if key == 'end' else R.side_bg
            data = fn(path, fx, fy)
        else:
            data = R.underlay_on(base, path, opacity, fx, fy)
        return _thumb(data, width)
    finally:
        R.set_theme('dark')


def item(style, path, kind='tile', width=360):
    """JPEG of a position tile (configuration page) or an option card with the photo, in the style's colours."""
    S = STYLES[style]
    R.set_theme(S['theme'])
    try:
        w, h = (2.3, 1.3) if kind == 'card' else (2.2, 1.6)
        if kind == 'card':
            im = R.card_photo(path, w, h)
        else:
            src = Image.open(R.asset(path))
            if src.mode in ('RGBA', 'LA', 'P') and src.convert('RGBA').getchannel('A').getextrema()[0] < 250:
                im = R.on_card(R.cutout_glow(R.alpha_crop(src.convert('RGBA'))), w, h, 0.9)
            else:
                im = R.grade(R.cover(path, w, h), sat=0.8, bright=0.92, warm=0.0)
        return _thumb(R.jpg(im), width)
    finally:
        R.set_theme('dark')
