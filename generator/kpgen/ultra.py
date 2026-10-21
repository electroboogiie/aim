# «Ultra» style (21.10.2026): KP for a crane in a kit (BMF 650 / 720 / 750), for НПЛ-10 + crane and for a truck
# converted to a timber carrier (КРАЗ / ЗІЛ + crane), built from scratch
# with python-pptx — dark graphite pages, rounded cards, large photos and type, BMF orange.
# The manufacturer drawings and the load chart with the crane are the ones the KP_2_v3 engine draws for the model
# (harvested from its build), so they stay as detailed as in the other KPs.
#   python3 -m kpgen.ultra --spec spec.json --out DIR [--no-pdf]
import argparse, io, json, math, os, re, sys
import numpy as np
from PIL import Image, ImageDraw, ImageFilter, ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from pptx.util import Pt

from . import raster as R
from .cache import Cache
from .pptx_util import EMU, I, by_id, by_name
from .profiles import PROFILES, DECAL_CONFIG, CARD_TWEAKS, OWN_PRODUCTION
from .spec import SpecError

W, H = 8.27, 11.69
BG, CARD, CARD2, LINE = (11, 12, 14), (23, 27, 33), (30, 34, 41), (44, 48, 56)
WHITE, MUTED, DIM, OR = (244, 244, 242), (156, 162, 172), (104, 110, 120), (242, 83, 10)
FONTS = os.path.join(R.ASSETS, 'fonts')
F_BLACK = os.path.join(FONTS, 'InterDisplay-Black.otf')
F_BOLD = os.path.join(FONTS, 'InterDisplay-Bold.otf')
TPL = os.path.normpath(os.path.join(os.path.dirname(os.path.abspath(__file__)), '..', 'templates'))
DPI = 300

# photos per crane: page «Кран» and the full-bleed last page (path, focus x, focus y)
PHOTOS = {
    'bmf720': dict(crane=('bmf/720_workshop.jpg', 0.5, 0.62), end=('bmf/720_b.jpg', 0.55, 0.45)),
    'bmf650': dict(crane=('bmf/650_b.jpg', 0.62, 0.5), end=('bmf/650_a.jpg', 0.32, 0.5)),
    'bmf750': dict(crane=('bmf/750_forest.webp', 0.5, 0.4), end=('bmf/750_forest_load.jpg', 0.5, 0.45)),
}
# НПЛ-10 kits: the user's photos and drawings (from their «НПЛ-10 Heavy Label» design and KPs)
NPL = dict(photo=('prod/npl10_tractor.png', 0.5, 0.55), end=('npl/npl10_workshop_hd.jpg', 0.45, 0.3),
           cutout='npl/npl10_crane_cutout.png', side='npl/npl10_dwg_side.png', rear='npl/npl10_dwg_rear.png')
# truck kits: the user's photos of the conversion (КП «КАМАЗ + 750») and what it includes
TRUCK_PH = dict(photo=('prod/truck_conversion.png', 0.5, 0.5), end=('prod/truck_conversion.png', 0.55, 0.5))
TRUCK_TILES = [('Бортова платформа', '{stakes} стійки під сортименти'), ('Рама {wall} мм', 'довжина {length} м'),
               ('Щит', 'передній щит платформи'), ('Піраміда з аутригерами', 'основа крана, стійкість під час роботи'),
               ('Гідробак у зборі', 'з оливою гідравлічною 90 л'), ('Насос аксіально-поршневий', 'привід гідросистеми крана'),
               ('Кран {model}', 'з ротатором {rotator} т і захватом {grapple} м²'), ('Монтаж', 'входить у вартість дообладнання')]
TILES = {'crane': None, 'base_outriggers': 'prod/tile_base.png', 'tank': 'prod/tile_tank.png', 'seat': 'prod/tile_seat.png'}


# ---------------------------------------------------------------- rasters
def font(path, px): return ImageFont.truetype(path, int(px))


def rtext(lines, path=F_BOLD, cap=120, pitch=1.22, track=-0.015):
    """Lines of (text, colour) segments drawn at cap height `cap` px. Returns (RGBA image, cap px, top offset px)."""
    probe = font(path, 100); hb = probe.getbbox('H'); size = 100 * cap / (hb[3] - hb[1]); f = font(path, size)
    tr = track * f.size
    wid = int(max(sum(f.getlength(c) + tr for t, _ in ln for c in t) for ln in lines) + f.size)
    im = Image.new('RGBA', (wid, int(cap * pitch * (len(lines) - 1) + f.size * 2.2)), (0, 0, 0, 0)); d = ImageDraw.Draw(im)
    base0 = f.size * 1.2
    for i, ln in enumerate(lines):
        x = f.size * 0.1
        for t, col in ln:
            for c in t:
                d.text((x, base0 + i * cap * pitch), c, font=f, anchor='ls', fill=tuple(col) + (255,)); x += f.getlength(c) + tr
    l, t, r, b = im.getbbox()
    return im.crop((l, t, r, b)), cap, base0 - cap - t


def png(im):
    b = io.BytesIO(); im.save(b, 'PNG', optimize=True); return b.getvalue()


def jpg(im, q=88):
    b = io.BytesIO(); im.convert('RGB').save(b, 'JPEG', quality=q); return b.getvalue()


def no_photo(label, w, h):
    """Card tile for a position without a photo: its volume or count in big orange type."""
    im = Image.new('RGB', (R.px(w), R.px(h)), CARD2)
    t = rtext([[(label, OR)]], F_BLACK, cap=int(im.height * 0.28))[0]
    im.paste(t, ((im.width - t.width) // 2, (im.height - t.height) // 2), t)
    return jpg(im)


def page_bg(glows, dots=True, dpi=120):
    """Graphite page with soft orange light; glows = [(cx, cy, r inches, alpha 0..1)]."""
    w, h = int(W * dpi), int(H * dpi)
    yy, xx = np.mgrid[0:h, 0:w].astype(np.float32) / dpi
    a = np.zeros((h, w), np.float32)
    for cx, cy, r, k in glows:
        d2 = ((xx - cx) ** 2 + (yy - cy) ** 2) / (r * r)
        a = np.maximum(a, k * np.exp(-d2 * 2.0))
    base = np.array(BG, np.float32); col = np.array(OR, np.float32)
    img = base[None, None, :] * (1 - a[..., None]) + col[None, None, :] * a[..., None]
    if dots:   # faint dot grid every 0.25 in
        step = int(0.25 * dpi); m = np.zeros((h, w), bool); m[step // 2::step, step // 2::step] = True
        img[m] += 14
    return jpg(Image.fromarray(np.clip(img, 0, 255).astype(np.uint8)), 90)


def crane_cut(blob, no):
    im = Image.open(io.BytesIO(blob)).convert('RGBA')
    if no != '720': im = R.redecal(im, DECAL_CONFIG, no)
    im.putalpha(im.getchannel('A').point(lambda v: 0 if v < 90 else v))
    return png(im.crop(im.getbbox()))


def floor_glow(w_in, h_in):
    """Orange light pool under the cover crane."""
    w, h = int(w_in * 150), int(h_in * 150)
    im = Image.new('RGBA', (w, h), OR + (0,)); m = Image.new('L', (w, h), 0)
    ImageDraw.Draw(m).ellipse((w * 0.08, h * 0.38, w * 0.92, h * 0.92), fill=150)
    im.putalpha(m.filter(ImageFilter.GaussianBlur(h * 0.16)))
    return png(im)


def photo(path, w, h, fx, fy, fade=None):
    im = R.grade(R.cover(path, w, h, fx, fy), sat=0.85, bright=0.92, contrast=1.08, warm=0.0).convert('RGB')
    if fade:   # fade into the page colour at the bottom (fraction of the height)
        ov = Image.new('RGBA', im.size, BG + (0,)); g = Image.new('L', (1, im.height))
        for y in range(im.height):
            t = max(0.0, (y / im.height - (1 - fade)) / fade); g.putpixel((0, y), int(255 * t * t * (3 - 2 * t)))
        ov.putalpha(g.resize(im.size)); im = Image.alpha_composite(im.convert('RGBA'), ov)
    return jpg(im)


def drawing_panel(parts, dims):
    """Manufacturer drawings (side view, folded view, plate) with their dimension labels baked in.
    parts: [(blob, x, y, w, h)] in inches; dims: [(text, cx, cy)]."""
    x0 = min(p[1] for p in parts) - 0.1; y0 = min(p[2] for p in parts) - 0.12
    x1 = max(p[1] + p[3] for p in parts) + 0.1; y1 = max(p[2] + p[4] for p in parts) + 0.12
    k = DPI
    im = Image.new('RGBA', (int((x1 - x0) * k), int((y1 - y0) * k)), (0, 0, 0, 0))
    for blob, x, y, w, h in parts:
        p = Image.open(io.BytesIO(blob)).convert('RGBA').resize((max(1, int(w * k)), max(1, int(h * k))), Image.LANCZOS)
        im.alpha_composite(p, (int((x - x0) * k), int((y - y0) * k)))
    d = ImageDraw.Draw(im); f = font(F_BOLD, 0.085 * k)
    for t, cx, cy in dims:
        l, tp, r, b = f.getbbox(t); tw, th = r - l, b - tp
        X, Y = (cx - x0) * k, (cy - y0) * k
        d.rounded_rectangle((X - tw / 2 - 10, Y - th / 2 - 8, X + tw / 2 + 10, Y + th / 2 + 8), radius=8, fill=CARD + (255,))
        d.text((X - tw / 2 - l, Y - th / 2 - tp), t, font=f, fill=WHITE + (255,))
    return png(im)


def load_chart(side, span, reach, l4, lf, gx=0.75, gw=6.9, axis_y=4.45):
    """Load-by-reach chart of the KP_2_v3 passport: the crane's side drawing on the metre axis (same geometry),
    bars at 4 m and at full reach with their kg. side = (blob, x, y, w, h) in inches."""
    blob, dx, dy, dw, dh = side
    x0, x1 = min(dx, gx - 0.25) - 0.05, max(dx + dw, gx + gw + 0.35) + 0.05
    y0, y1 = min(dy, axis_y - 1.9) - 0.1, axis_y + 0.5
    k = DPI; X = lambda v: (v - x0) * k; Y = lambda v: (v - y0) * k
    im = Image.new('RGBA', (int((x1 - x0) * k), int((y1 - y0) * k)), (0, 0, 0, 0)); d = ImageDraw.Draw(im)
    fm = font(F_BOLD, 0.115 * k)
    for m in range(int(span) + 1):
        if m > reach + 0.4: break
        x = X(gx + gw * m / span)
        d.line((x, Y(axis_y - 1.75), x, Y(axis_y)), fill=LINE + (255,), width=3)
        t = f'{m} м'; l, tp, r, b = fm.getbbox(t); d.text((x - (r - l) / 2 - l, Y(axis_y + 0.12) - tp), t, font=fm, fill=DIM + (255,))
    d.line((X(gx), Y(axis_y), X(gx + gw * min(span, reach + 0.3) / span), Y(axis_y)), fill=MUTED + (255,), width=5)
    p = Image.open(io.BytesIO(blob)).convert('RGBA').resize((int(dw * k), int(dh * k)), Image.LANCZOS)
    im.alpha_composite(p, (int(X(dx)), int(Y(dy))))
    ky = min(1.2 / 1100, 1.3 / max(l4, lf)); bw = 0.46
    fk, fs = font(F_BLACK, 0.22 * k), font(F_BOLD, 0.11 * k)
    for m, kg, col, sub in ((4, l4, OR, 'на 4 м'), (reach, lf, WHITE, f"на {str(reach).replace('.', ',')} м")):
        cx = X(gx + gw * m / span); top = Y(axis_y - kg * ky)
        d.rounded_rectangle((cx - bw * k / 2, top, cx + bw * k / 2, Y(axis_y) - 2), radius=10, fill=col + (235,))
        t = f'{kg} кг'; l, tp, r, b = fk.getbbox(t)
        tx = min(max(cx - (r - l) / 2, 0), im.width - (r - l) - 4)
        d.text((tx - l, top - 0.46 * k - tp), t, font=fk, fill=WHITE + (255,))
        l2, t2, r2, b2 = fs.getbbox(sub); d.text((min(max(cx - (r2 - l2) / 2, 0), im.width - (r2 - l2) - 4) - l2, top - 0.17 * k - t2), sub, font=fs, fill=MUTED + (255,))
    return png(im)


def v3_pieces(spec, build_dir):
    """Build the KP in the KP_2_v3 engine and take its crane drawings: (engine, [side, folded, plate] pictures,
    dimension labels, load-chart side drawing). Pictures are (blob, x, y, w, h) in inches."""
    from .build import Engine
    R.set_theme('dark')
    e = Engine(dict(spec, style='dark', overrides={}, photos={}), build_dir); e.build()
    an, pp = e.pages['anatomy'], e.pages['passport']
    g = lambda sh: (sh.image.blob, sh.left / EMU, sh.top / EMU, sh.width / EMU, sh.height / EMU)
    parts = [g(by_id(an, i)) for i in (85, 88, 82)]
    dims = [(sh.text_frame.text, (sh.left + sh.width / 2) / EMU, (sh.top + sh.height / 2) / EMU)
            for sh in an.shapes if sh.name.startswith('dim_') and sh.text_frame.text.strip()]
    side = g(by_name(pp, 'crane_side_drawing')[0])
    R.set_theme('dark')
    return e, parts, dims, side


# ---------------------------------------------------------------- pptx helpers
def rgb(c): return RGBColor(*c)


class Deck:
    def __init__(self):
        self.prs = Presentation(); self.prs.slide_width, self.prs.slide_height = I(W), I(H)

    def slide(self, bg=None):
        s = self.prs.slides.add_slide(self.prs.slide_layouts[6])
        if bg: self.img(s, bg, 0, 0, W, H)
        return s

    @staticmethod
    def box(s, x, y, w, h, fill=CARD, line=None, r=0.16, lw=0.75):
        sh = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if r else MSO_SHAPE.RECTANGLE, I(x), I(y), I(w), I(h))
        if r: sh.adjustments[0] = min(0.5, r / min(w, h))
        if fill is None: sh.fill.background()
        else: sh.fill.solid(); sh.fill.fore_color.rgb = rgb(fill)
        if line: sh.line.color.rgb = rgb(line); sh.line.width = Pt(lw)
        else: sh.line.fill.background()
        sh.shadow.inherit = False
        return sh

    @staticmethod
    def text(s, x, y, w, h, paras, size=11, color=WHITE, bold=False, align='l', anchor='t', spacing=1.1, font='Arial', track=None):
        tb = s.shapes.add_textbox(I(x), I(y), I(w), I(h)); tf = tb.text_frame; tf.word_wrap = True
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = {'t': MSO_ANCHOR.TOP, 'm': MSO_ANCHOR.MIDDLE, 'b': MSO_ANCHOR.BOTTOM}[anchor]
        if isinstance(paras, str): paras = [paras]
        for i, para in enumerate(paras):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = {'l': PP_ALIGN.LEFT, 'c': PP_ALIGN.CENTER, 'r': PP_ALIGN.RIGHT}[align]; p.line_spacing = spacing
            runs = para if isinstance(para, list) else [(para, {})]
            for t, o in runs:
                r = p.add_run(); r.text = t; f = r.font
                f.size = Pt(o.get('size', size)); f.bold = o.get('bold', bold); f.name = o.get('font', font)
                f.color.rgb = rgb(o.get('color', color))
                if o.get('track', track) is not None: r._r.get_or_add_rPr().set('spc', str(int(o.get('track', track) * 100)))
        return tb

    @staticmethod
    def img(s, data, x, y, w=None, h=None, r=None):
        pic = s.shapes.add_picture(io.BytesIO(data), I(x), I(y), I(w) if w else None, I(h) if h else None)
        if r:
            geom = pic._element.spPr.find(qn('a:prstGeom')); geom.set('prst', 'roundRect')
            av = geom.find(qn('a:avLst'))
            if av is None: av = geom.makeelement(qn('a:avLst'), {}); geom.append(av)
            gd = av.makeelement(qn('a:gd'), {'name': 'adj', 'fmla': f'val {int(min(50000, r / min(pic.width, pic.height) * EMU * 100000))}'})
            av.append(gd)
        return pic

    def rt(self, s, data_cap, x, y, cap_in, maxw=None):
        """Place a raster headline: cap height cap_in inches, the tops of the capitals at y."""
        data, cap, off = data_cap; im = Image.open(io.BytesIO(data))
        k = cap_in / cap
        if maxw and im.width * k > maxw: k = maxw / im.width
        return self.img(s, data, x, y - off * k, im.width * k, im.height * k)


# ---------------------------------------------------------------- the KP
class Ultra:
    def __init__(self, spec, build_dir):
        if spec['kit_type'] not in ('crane_kit', 'npl_crane', 'truck_crane'):
            raise SpecError('Стиль «Ultra» є для крана в комплекті, НПЛ-10 + кран і вантажівки з краном. Оберіть інший стиль.')
        self.npl = spec['kit_type'] == 'npl_crane'; self.truck = spec['kit_type'] == 'truck_crane'
        # truck brand from the conversion line («Дообладнання автомобіля ЗІЛ у лісовоз») or the KP title («ЗІЛ + BMF 650»)
        conv = next((r['title'] for r in spec['rows'] if r.get('kind') == 'service'), '')
        mt = re.search(r'автомобіля\s+(\S+)', conv)
        self.truck_name = (mt.group(1) if mt else spec['title'].split(' + ')[0].strip()) if self.truck else ''
        self.spec = spec; self.c = spec['crane']; self.cs = self.c['specs']
        self.prof = PROFILES[self.c['profile']]; self.no = str(self.prof.get('model_no') or '')
        self.build_dir = build_dir; self.cache = Cache(build_dir)
        self.m = self.c['model']

    def raster(self, kind, fn, *args, assets=()): return self.cache.get('ultra-' + kind, fn, args, assets)

    def photo_of(self, key, d):
        """(path, fx, fy) of a photo page: the settings' photo when one is set, else the style's."""
        o = ((self.spec.get('overrides') or {}).get('bg') or {}).get(key) or {}
        if o.get('path') and R.asset_ok(o['path']): return o['path'], float(o.get('fx', 0.5)), float(o.get('fy', 0.45))
        return d

    def headline(self, lines, path=F_BOLD):
        lines = [[(t, tuple(c)) for t, c in ln] for ln in lines]
        res = self.cache.get('ultra-hl', lambda L, p: (lambda r: (png(r[0]), [r[1], r[2]]))(rtext(L, p)), (lines, path))
        data, (cap, off) = res
        return data, cap, off

    # -- pieces the KP_2_v3 engine draws for this crane
    def harvest(self):
        e, parts, dims, side = v3_pieces(self.spec, self.build_dir)
        cut = by_id(list(Presentation(os.path.join(TPL, 'KP_2_v3.pptx')).slides)[4], 14).image.blob
        logo = by_id(list(Presentation(os.path.join(TPL, 'KP_2_v3.pptx')).slides)[1], 3).image.blob
        d6 = list(Presentation(os.path.join(TPL, 'KP_npl_label.pptx')).slides)[0]
        qrs = [by_name(d6, n)[0].image.blob for n in ('qr_сайт', 'qr_tiktok', 'qr_instagram')]
        R.set_theme('dark')
        self.td = e.tdata() if self.npl else None
        return parts, dims, side, e.span, cut, logo, qrs

    def build(self):
        sp, cs = self.spec, self.cs
        parts, dims, side, span, cut, logo, qrs = self.harvest()
        self.logo = logo; self.qrs = qrs
        self.cut = self.raster('cut', crane_cut, cut, self.no)
        self.panel = self.raster('panel', drawing_panel, parts, dims)
        reach = float(cs['reach_m'].replace(',', '.'))
        self.chart = self.raster('chart', load_chart, side, span, reach, int(cs['lift_4m_kg']), int(cs['lift_full_kg']))
        D = self.deck = Deck()
        self.pages = []
        order = ([self.p_cover, self.p_npl, self.p_npl_spec] if self.npl else [self.p_cover, self.p_truck] if self.truck else [self.p_cover]) + \
            [self.p_crane, self.p_drawings, self.p_chart, self.p_kit, self.p_options, self.p_price] + \
            ([self.p_own] if self.npl or self.truck else []) + [self.p_end]
        for fn in order: fn()
        self.chrome()
        return D.prs

    # header / footer on every page but the cover
    def chrome(self):
        D, sp = self.deck, self.spec; n = len(self.pages)
        for i, (s, dark_top) in enumerate(self.pages, 1):
            if i == 1: continue
            D.img(s, self.logo, 0.5, 0.42, 0.62)
            D.text(s, 1.35, 0.44, 4.2, 0.2, f"НПЛ-10  ×  {self.m.upper()}" if self.npl else f"{self.truck_name}  ×  {self.m.upper()}  ·  ЛІСОВОЗ" if self.truck else f"{self.m.upper()}  ·  КРАН У КОМПЛЕКТІ", 7.5, MUTED, True, track=1.5)
            D.text(s, 5.2, 0.44, 2.57, 0.2, f"КП № {sp['number']}  ·  {sp['date']}", 7.5, MUTED, True, 'r', track=1.0)
            D.box(s, 0.5, 11.02, 7.27, 0.012, LINE, r=0)
            site = (sp['contacts']['sites'] or [''])[0]; ph = (sp['contacts']['phones'] or [''])[0]
            D.text(s, 0.5, 11.14, 5, 0.2, f"{site.upper()}   ·   {ph}", 7, DIM, True, track=1.0)
            D.text(s, 6.27, 11.1, 1.5, 0.25, [[(f'{i:02d}', dict(color=WHITE, size=10)), (f' / {n:02d}', dict(color=DIM, size=8))]], 10, WHITE, True, 'r')

    def kicker(self, s, x, y, t):
        t = f'{len(self.pages) - 1:02d} · {t}'   # section number = page number after the cover
        self.deck.box(s, x, y + 0.035, 0.09, 0.09, OR, r=0.045)
        self.deck.text(s, x + 0.18, y, 6, 0.2, t.upper(), 7.5, OR, True, track=2.0)

    # 01 cover
    def p_cover(self):
        D, sp, cs = self.deck, self.spec, self.cs
        s = D.slide(self.raster('bg-cover', page_bg, ((4.6, 6.4, 4.2, 0.5), (8.4, 0.2, 3.0, 0.16))))
        self.pages.append((s, True))
        D.img(s, self.logo, 0.5, 0.45, 0.85)
        D.text(s, 3.5, 0.48, 4.27, 0.22, f"КОМЕРЦІЙНА ПРОПОЗИЦІЯ № {sp['number']}", 8, MUTED, True, 'r', track=2.0)
        D.text(s, 3.5, 0.7, 4.27, 0.22, sp['date'], 8, DIM, True, 'r', track=2.0)
        if self.truck:
            D.rt(s, self.headline([[(self.truck_name + ' ', WHITE), ('+ ' + (self.no or self.m), OR)]], F_BLACK), 0.47, 1.4, 1.45, maxw=7.3)
            D.text(s, 0.5, 3.2, 7.27, 0.4, [[('Лісовоз з краном ', {}), (self.m, dict(color=OR))]], 24, WHITE, True)
            D.text(s, 0.5, 3.72, 7.27, 0.3, f"дообладнання вантажівки {self.truck_name} під ключ · ротатор {cs['rotator_t']} т і грейфер {cs['grapple_m2']} м² у комплекті", 12, MUTED)
            d = self.photo_of('truck', TRUCK_PH['photo'])
            D.img(s, self.raster('photo', photo, d[0], 7.27, 4.25, d[1], d[2], assets=(d[0],)), 0.5, 4.2, 7.27, 4.25, r=0.22)
            kpis = [(self.truck_spec('stakes', '4'), 'шт', 'стійки платформи'), (cs['reach_m'], 'м', f'виліт крана {self.m}'),
                    (cs['torque_knm'], 'кНм', 'підйомний момент')]
        elif self.npl:
            D.rt(s, self.headline([[('НПЛ-10', WHITE)]], F_BLACK), 0.47, 1.4, 1.45, maxw=7.3)
            D.text(s, 0.5, 3.2, 7.27, 0.4, [[('Лісовоз з краном ', {}), (self.m, dict(color=OR))]], 24, WHITE, True)
            D.text(s, 0.5, 3.72, 7.27, 0.3, f"напівпричіп власного виробництва · ротатор {cs['rotator_t']} т і грейфер {cs['grapple_m2']} м² у комплекті", 12, MUTED)
            D.img(s, self.raster('floor', floor_glow, 8.0, 2.8), 0.13, 6.1, 8.0, 2.8)
            cut = self.raster('npl-cut', lambda p_: png(R.alpha_crop(Image.open(R.asset(p_)).convert('RGBA'))), NPL['cutout'], assets=(NPL['cutout'],))
            im = Image.open(io.BytesIO(cut)); h = 4.95; w = h * im.width / im.height
            D.img(s, cut, (W - w) / 2 + 0.2, 9.0 - h, w, h)
            kpis = [('10 000', 'кг', 'вантажопідйомність НПЛ-10'), (cs['reach_m'], 'м', f'виліт крана {self.m}'),
                    (cs['torque_knm'], 'кНм', 'підйомний момент')]
        else:
            D.rt(s, self.headline([[('BMF ', WHITE), (self.no or self.m, OR)]], F_BLACK), 0.47, 1.4, 1.45, maxw=7.3)
            D.text(s, 0.5, 3.2, 7.27, 0.4, 'Кран-маніпулятор у комплекті', 24, WHITE, True)
            D.text(s, 0.5, 3.72, 7.27, 0.3, f"з ротатором {cs['rotator_t']} т і грейфером {cs['grapple_m2']} м² — уже в ціні крана", 12, MUTED)
            D.img(s, self.raster('floor', floor_glow, 8.0, 2.8), 0.13, 6.1, 8.0, 2.8)
            im = Image.open(io.BytesIO(self.cut)); w = 8.0; h = w * im.height / im.width
            D.img(s, self.cut, (W - w) / 2, 8.45 - h, w, h)
            kpis = [(cs['reach_m'], 'м', 'максимальний виліт'), (cs['torque_knm'], 'кНм', 'підйомний момент'),
                    (cs['lift_4m_kg'], 'кг', 'вантаж на вильоті 4 м')]
        cw = (7.27 - 0.3) / 3
        for i, (v, u, l) in enumerate(kpis):
            x = 0.5 + i * (cw + 0.15)
            D.box(s, x, 8.65, cw, 1.2, CARD, LINE)
            D.text(s, x + 0.25, 8.83, cw - 0.4, 0.5, [[(v, dict(size=28)), (f'  {u}', dict(size=12, color=OR))]], 28, WHITE, True)
            D.text(s, x + 0.25, 9.43, cw - 0.4, 0.3, l, 9.5, MUTED)
        client = sp['client']['name']
        D.box(s, 0.5, 10.3, 7.27, 0.012, LINE, r=0)
        D.text(s, 0.5, 10.45, 3.6, 0.5, [[('ДЛЯ  ', dict(color=DIM, size=7.5, track=2.0)), ((client or '—'), dict(size=11))]], 11, WHITE, True)
        ct = sp['contacts']
        D.text(s, 4.1, 10.45, 3.67, 0.6, [' · '.join(ct['phones'][:2]), ' · '.join(ct['sites'][:2])], 9, MUTED, True, 'r', spacing=1.25)
        D.text(s, 0.5, 11.2, 7.27, 0.2, sp['maker']['line'].upper(), 6.5, DIM, True, track=1.5)

    def truck_spec(self, k, d):
        r = next((r for r in self.spec['rows'] if r.get('kind') == 'service'), {})
        m = {'stakes': r'(\d+)\s*стійк', 'wall': r'рама\s*(\d+)\s*мм', 'length': r'/\s*([\d,]+)\s*м'}[k]
        mm = re.search(m, r.get('desc') or ''); return mm.group(1) if mm else d

    # truck: what the conversion includes
    def p_truck(self):
        D, cs, t = self.deck, self.cs, self.truck_name
        s = D.slide(self.raster('bg-a', page_bg, ((8.3, 6.5, 3.2, 0.14),)))
        self.pages.append((s, False))
        d = self.photo_of('npl', None)
        if d:
            D.img(s, self.raster('photo', photo, d[0], 7.27, 4.2, d[1], d[2], assets=(d[0],)), 0.5, 0.95, 7.27, 4.2, r=0.2)
        else:   # the user's scheme of the conversion (platform, pyramid with outriggers, tank) on a card
            sch = 'prod/truck_pyramid.png'
            D.box(s, 0.5, 0.95, 7.27, 4.2, CARD, LINE, r=0.2)
            D.img(s, self.raster('scheme', lambda p_, w_, h_: jpg(R.on_card(Image.open(R.asset(p_)).convert('RGBA'), w_, h_, 0.95)), sch, 7.0, 3.95, assets=(sch,)), 0.63, 1.08, 7.0, 3.95, r=0.16)
            D.text(s, 5.0, 4.75, 2.6, 0.2, '* схематичне зображення', 7.5, DIM, align='r')
        D.box(s, 0.75, 4.67, 2.3, 0.3, OR, r=0.15)
        D.text(s, 0.75, 4.67, 2.3, 0.3, 'СХЕМА ДООБЛАДНАННЯ' if not d else 'ДООБЛАДНАННЯ', 8, WHITE, True, 'c', 'm', track=0.5)
        self.kicker(s, 0.5, 5.45, 'дообладнання · власне виробництво')
        D.rt(s, self.headline([[('Лісовоз на базі', WHITE)], [(t, OR)]]), 0.48, 5.8, 0.42)
        conv = next((r for r in self.spec['rows'] if r.get('kind') == 'service'), {})
        D.text(s, 0.5, 6.87, 7.27, 0.9, [f"Переобладнуємо вантажівку {t} у лісовоз для перевезення колод і брусів: бортова платформа зі стійками, "
                                         f"піраміда з аутригерами під кран-маніпулятор {self.m}, гідросистема і монтаж.",
                                         (conv.get('desc') or '').capitalize()], 10.5, (210, 213, 218), spacing=1.25)
        fm = dict(stakes=self.truck_spec('stakes', '4'), wall=self.truck_spec('wall', '6'), length=self.truck_spec('length', '5'),
                  model=self.m, rotator=cs['rotator_t'], grapple=cs['grapple_m2'])
        tw, gap = (7.27 - 0.15) / 2, 0.12; rh = min(0.68, (10.85 - 7.85 - gap * 3) / 4)
        for i, (a, b) in enumerate(TRUCK_TILES):
            x, y = 0.5 + (i % 2) * (tw + 0.15), 7.85 + (i // 2) * (rh + gap)
            D.box(s, x, y, tw, rh, CARD, LINE, r=0.14)
            D.text(s, x + 0.18, y, 0.35, rh, f'{i + 1:02d}', 8.5, OR, True, anchor='m')
            D.text(s, x + 0.52, y + 0.04, tw - 0.62, rh - 0.08, [[(a.format(**fm), dict(bold=True, size=10.5, color=WHITE))], [(b.format(**fm), dict(size=8, color=MUTED))]], 10.5, WHITE, anchor='m', spacing=1.05)

    # НПЛ-10: photo, description, features
    def p_npl(self):
        D, td = self.deck, self.td
        s = D.slide(self.raster('bg-a', page_bg, ((8.3, 6.5, 3.2, 0.14),)))
        self.pages.append((s, False))
        d = self.photo_of('npl', NPL['photo'])
        D.img(s, self.raster('photo', photo, d[0], 7.27, 4.2, d[1], d[2], assets=(d[0],)), 0.5, 0.95, 7.27, 4.2, r=0.2)
        D.box(s, 0.75, 4.67, 1.05, 0.3, OR, r=0.15)
        D.text(s, 0.75, 4.67, 1.05, 0.3, 'НПЛ-10', 9, WHITE, True, 'c', 'm', track=1.0)
        self.kicker(s, 0.5, 5.45, 'напівпричіп · власне виробництво')
        D.rt(s, self.headline([[('Напівпричіп-лісовоз', WHITE)], [('НПЛ-10', OR)]]), 0.48, 5.8, 0.42)
        D.text(s, 0.5, 6.87, 7.27, 0.9, [td['lead'], td['lead2']], 10.5, (210, 213, 218), spacing=1.25)
        tw, gap = (7.27 - 0.15) / 2, 0.12
        tiles = td['tiles'][:8]; rh = min(0.68, (10.85 - 7.85 - gap * 3) / 4)
        for i, (t, sub) in enumerate(tiles):
            x, y = 0.5 + (i % 2) * (tw + 0.15), 7.85 + (i // 2) * (rh + gap)
            D.box(s, x, y, tw, rh, CARD, LINE, r=0.14)
            D.text(s, x + 0.18, y, 0.35, rh, f'{i + 1:02d}', 8.5, OR, True, anchor='m')
            D.text(s, x + 0.52, y + 0.04, tw - 0.62, rh - 0.08, [[(t, dict(bold=True, size=10.5, color=WHITE))], [(sub, dict(size=8, color=MUTED))]], 10.5, WHITE, anchor='m', spacing=1.05)

    # НПЛ-10: drawings + specification
    def p_npl_spec(self):
        D, td = self.deck, self.td
        s = D.slide(self.raster('bg-b', page_bg, ((-0.2, 3.2, 3.2, 0.14),)))
        self.pages.append((s, False))
        self.kicker(s, 0.5, 0.95, 'габарити та характеристики · НПЛ-10')
        D.rt(s, self.headline([[('Міцний', WHITE)], [('на 10 тонн', OR)]]), 0.48, 1.3, 0.42)
        bw, bh, top = 7.27, 3.35, 2.55
        D.box(s, 0.5, top, bw, bh, CARD, LINE, r=0.2)
        D.text(s, 0.75, top + 0.17, 5, 0.2, 'ГАБАРИТНІ РОЗМІРИ, ММ*', 7.5, MUTED, True, track=1.5)
        side = self.raster('npl-dwg', lambda p_: png(R.alpha_crop(Image.open(R.asset(p_)).convert('RGBA'))), NPL['side'], assets=(NPL['side'],))
        rear = self.raster('npl-dwg', lambda p_: png(R.alpha_crop(Image.open(R.asset(p_)).convert('RGBA'))), NPL['rear'], assets=(NPL['rear'],))
        for data, x, w in ((side, 0.75, 4.45), (rear, 5.45, 2.1)):
            im = Image.open(io.BytesIO(data)); k = min(w / im.width, (bh - 0.75) / im.height)
            D.img(s, data, x + (w - im.width * k) / 2, top + 0.5 + (bh - 0.65 - im.height * k) / 2, im.width * k, im.height * k)
        rows = td['spec_rows']; y0 = top + bh + 0.25; rh = min(0.3, (10.6 - y0) / len(rows))
        for i, (l, v, u) in enumerate(rows):
            y = y0 + i * rh
            if i % 2 == 0: D.box(s, 0.5, y, 7.27, rh, CARD, r=0.08)
            D.text(s, 0.75, y, 5, rh, l, 10, (214, 217, 222), anchor='m')
            D.text(s, 5.0, y, 2.15, rh, [[(v, dict(bold=True, color=WHITE)), ((' ' + u) if u else '', dict(color=MUTED, size=9))]], 10.5, WHITE, align='r', anchor='m')
        D.text(s, 0.5, y0 + len(rows) * rh + 0.06, 7.27, 0.3, td['note'], 7.5, DIM, spacing=1.1)

    # НПЛ-10 kits: own production (photos from the user's KPs)
    def p_own(self):
        D, O = self.deck, OWN_PRODUCTION
        s = D.slide(self.raster('bg-c', page_bg, ((8.4, 2.4, 3.6, 0.2), (-0.4, 10.5, 3.0, 0.1))))
        self.pages.append((s, False))
        self.kicker(s, 0.5, 0.95, 'власне виробництво')
        a, b = O['headline']
        facts = O['facts'] if not self.truck else [('до 60', 'ДНІВ ПІД ЗАМОВЛЕННЯ'), ('під ключ', 'ПЛАТФОРМА, ПІРАМІДА, ГІДРАВЛІКА'),
                                                      ('монтаж', 'ВХОДИТЬ У ВАРТІСТЬ'), ('BMF', 'ОФІЦІЙНИЙ ДИСТРИБ’ЮТОР')]
        D.rt(s, self.headline([[(a.rstrip('.').capitalize(), WHITE)], [(b.rstrip('.').capitalize(), OR)]]), 0.48, 1.3, 0.42)
        D.text(s, 0.5, 2.5, 7.27, 0.75, O['quote'], 11, (214, 217, 222), spacing=1.25)
        n = len(facts); gap = 0.12; cw = (7.27 - gap * (n - 1)) / n
        for i, (v, cap) in enumerate(facts):
            x = 0.5 + i * (cw + gap)
            D.box(s, x, 3.45, cw, 1.15, CARD, LINE, r=0.14)
            D.text(s, x + 0.15, 3.58, cw - 0.25, 0.45, v, 19, OR if i == 0 else WHITE, True)
            D.text(s, x + 0.15, 4.05, cw - 0.25, 0.5, cap.capitalize(), 8, MUTED, spacing=1.05)
        photos = [p_ for _, p_, _, _ in O['photos']][:6]
        if self.truck: photos = sorted(photos, key=lambda p_: 'truck' not in p_)
        foc = {p_: (fx, fy) for _, p_, fx, fy in O['photos']}
        cw, ch, gap = (7.27 - 0.15) / 2, 1.95, 0.15
        for i, p_ in enumerate(photos):
            x, y = 0.5 + (i % 2) * (cw + gap), 4.85 + (i // 2) * (ch + gap)
            D.img(s, self.raster('photo', photo, p_, cw, ch, *foc[p_], assets=(p_,)), x, y, cw, ch, r=0.16)

    # the crane
    def p_crane(self):
        D, cs = self.deck, self.cs
        s = D.slide(self.raster('bg-a', page_bg, ((8.3, 6.5, 3.2, 0.14),)))
        self.pages.append((s, False))
        d = self.photo_of('crane', PHOTOS[self.c['profile']]['crane'])
        D.img(s, self.raster('photo', photo, d[0], 7.27, 4.7, d[1], d[2], assets=(d[0],)), 0.5, 0.95, 7.27, 4.7, r=0.2)
        D.box(s, 0.75, 5.17, 1.05, 0.3, OR, r=0.15)
        D.text(s, 0.75, 5.17, 1.05, 0.3, self.m, 9, WHITE, True, 'c', 'm', track=1.0)
        self.kicker(s, 0.5, 6.0, 'кран-маніпулятор')
        D.rt(s, self.headline([[('Лісовий кран', WHITE)], [(self.m, OR)]]), 0.48, 6.35, 0.42)
        lead = self.c.get('lead') or self.prof.get('lead'); lead2 = self.c.get('lead2') or self.prof.get('lead2')
        D.text(s, 0.5, 7.42, 7.27, 0.95, [lead, lead2], 10.5, (210, 213, 218), spacing=1.25)
        yes = lambda k: (cs.get(k) or '+') in ('так', '+')
        feats = [('Сталь SSAB Strenx', 'стріла з високоміцної сталі'), ('Триточкове навішування', 'на причіп або напряму на трактор'),
                 ('Ширококутний шарнір' if yes('wide_angle') else 'Поворотна колона', f"поворот стріли {cs['turn_deg']}°"),
                 ('Телескоп' if yes('telescope') else 'Підйомний момент', f"виліт до {cs['reach_m']} м" if yes('telescope') else f"{cs['torque_knm']} кНм")]
        tw = (7.27 - 0.15) / 2
        for i, (t, sub) in enumerate(feats):
            x, y = 0.5 + (i % 2) * (tw + 0.15), 8.45 + (i // 2) * 1.2
            D.box(s, x, y, tw, 1.05, CARD, LINE)
            D.text(s, x + 0.25, y + 0.2, 0.5, 0.25, f'0{i + 1}', 9, OR, True)
            D.text(s, x + 0.75, y + 0.17, tw - 1.0, 0.35, t, 13, WHITE, True)
            D.text(s, x + 0.75, y + 0.55, tw - 1.0, 0.35, sub, 10, MUTED)

    # 03 manufacturer drawings
    def p_drawings(self):
        D, cs = self.deck, self.cs
        s = D.slide(self.raster('bg-b', page_bg, ((-0.2, 3.2, 3.2, 0.14),)))
        self.pages.append((s, False))
        self.kicker(s, 0.5, 0.95, 'креслення виробника')
        D.rt(s, self.headline([[('Робоча зона', WHITE)], [('і габарити, мм', OR)]]), 0.48, 1.3, 0.42)
        im = Image.open(io.BytesIO(self.panel)); bw, bh = 7.27, 6.2
        D.box(s, 0.5, 2.55, bw, bh, CARD, LINE, r=0.2)
        k = min((bw - 0.5) / im.width, (bh - 0.5) / im.height); w, h = im.width * k, im.height * k
        D.img(s, self.panel, 0.5 + (bw - w) / 2, 2.55 + (bh - h) / 2, w, h)
        kp = [(cs['reach_m'], 'м', 'максимальний виліт'), (cs['lift_4m_kg'], 'кг', 'на вильоті 4 м*'),
              (cs['lift_full_kg'], 'кг', 'на повному вильоті*'), (cs['turn_deg'], '°', 'поворот стріли')]
        cw = (7.27 - 0.45) / 4
        for i, (v, u, l) in enumerate(kp):
            x = 0.5 + i * (cw + 0.15)
            D.box(s, x, 9.0, cw, 1.15, CARD, LINE)
            D.text(s, x + 0.2, 9.17, cw - 0.3, 0.45, [[(v, dict(size=22)), (f' {u}', dict(size=10, color=OR))]], 22, WHITE, True)
            D.text(s, x + 0.2, 9.68, cw - 0.3, 0.3, l, 9, MUTED)
        D.text(s, 0.5, 10.4, 7.27, 0.3, '* без захвату і ротатора. Креслення й дані — виробник BMF (bmf.ee).', 8, DIM)

    # 04 load chart + full specification
    def p_chart(self):
        D, cs = self.deck, self.cs
        s = D.slide(self.raster('bg-a', page_bg, ((8.3, 6.5, 3.2, 0.14),)))
        self.pages.append((s, False))
        self.kicker(s, 0.5, 0.95, 'технічний паспорт')
        D.rt(s, self.headline([[('Вантажопідйомність', WHITE)], [('за вильотом', OR)]]), 0.48, 1.3, 0.42)
        im = Image.open(io.BytesIO(self.chart)); bw, bh = 7.27, 3.3
        D.box(s, 0.5, 2.55, bw, bh, CARD, LINE, r=0.2)
        D.text(s, 0.75, 2.72, 5, 0.2, 'ВАНТАЖ БЕЗ ЗАХВАТУ І РОТАТОРА, КГ', 7.5, MUTED, True, track=1.5)
        k = min((bw - 0.3) / im.width, (bh - 0.5) / im.height); w, h = im.width * k, im.height * k
        D.img(s, self.chart, 0.5 + (bw - w) / 2, 2.55 + bh - h - 0.12, w, h)
        g = lambda k_, d='—': cs.get(k_) or d
        flow = g('valve_flow_l'); flow = '' if flow == '—' else f" ({flow} л)"
        rows = [('Максимальний виліт', cs['reach_m'], 'м'), ('Підйомний момент, брутто', cs['torque_knm'], 'кНм'),
                ('Вантажопідйомність на 4 м*', cs['lift_4m_kg'], 'кг'), ('Підйомна сила на повному вильоті*', cs['lift_full_kg'], 'кг'),
                ('Захват (у комплекті)', cs['grapple_m2'], 'м²'), ('Ротатор (у комплекті)', cs['rotator_t'], 'т'),
                ('Гідравлічний клапан, стандарт', g('valve') + flow, ''), ('Поворотний момент', g('slew_knm'), 'кНм'),
                ('Циліндри повороту', g('slew_cyl'), 'шт'), ('Поворот стріли', cs['turn_deg'], '°'), ('Робочий тиск', g('pressure_bar'), 'бар'),
                ('Продуктивність гідронасоса', g('pump_lpm'), 'л/хв'), ('Телескоп', 'так' if g('telescope', '+') in ('так', '+') else 'ні', ''),
                ('Тип опорних лап', g('legs'), ''), ('Вага без опори, захвату, ротатора', g('weight_kg'), 'кг')]
        y0, rh = 6.1, 0.315
        for i, (l, v, u) in enumerate(rows):
            y = y0 + i * rh
            if i % 2 == 0: D.box(s, 0.5, y, 7.27, rh, CARD, r=0.08)
            D.text(s, 0.75, y, 5, rh, l, 10, (214, 217, 222), anchor='m')
            D.text(s, 5.0, y, 2.15, rh, [[(v, dict(bold=True, color=WHITE)), ((' ' + u) if u else '', dict(color=MUTED, size=9))]], 10.5, WHITE, align='r', anchor='m')
        D.text(s, 0.5, y0 + len(rows) * rh + 0.08, 7.27, 0.25, '* без захвату і ротатора. Дані виробника BMF.', 8, DIM)

    # 05 kit composition
    def p_kit(self):
        D, sp, cs = self.deck, self.spec, self.cs
        s = D.slide(self.raster('bg-b', page_bg, ((-0.2, 3.2, 3.2, 0.14),)))
        self.pages.append((s, False))
        self.kicker(s, 0.5, 0.95, 'склад комплекту')
        D.rt(s, self.headline([[('Що входить', WHITE)], [('у комплект', OR)]]), 0.48, 1.3, 0.42)
        rows = sp['rows']; n = len(rows)
        top, bottom, gap = 2.55, 9.35, 0.15
        ch = min(2.6 if n <= 2 else 1.55, (bottom - top - gap * (n - 1)) / n); big = n <= 2
        for i, r in enumerate(rows):
            x, y, cw = 0.5, top + i * (ch + gap), 7.27
            D.box(s, x, y, cw, ch, CARD, LINE, r=0.18)
            ph = ch - 0.24; pw = min(ph * 1.6, 2.7)
            D.img(s, self.tile(r, pw, ph), x + 0.12, y + 0.12, pw, ph, r=0.12)
            tx = x + 0.12 + pw + 0.3
            pt = r['price_text'] or 'у комплекті'; pc = OR if any(c_.isdigit() for c_ in pt) else MUTED
            if big:   # one or two positions: title, description and price stacked in a wide column
                tw = cw - (tx - x) - 0.3
                D.text(s, tx, y + 0.3, tw, 0.22, f'{i + 1:02d}', 8.5, OR, True, track=1.0)
                body = [[(r['title'], dict(bold=True, size=18))]]
                if r['desc']: body.append([(r['desc'], dict(size=11, color=MUTED))])
                D.text(s, tx, y + 0.6, tw, ch - 1.3, body, 18, WHITE, spacing=1.15)
                D.text(s, tx, y + ch - 0.7, tw, 0.45, pt, 20, pc, True, anchor='b')
                continue
            prw = 1.75; tw = cw - (tx - x) - prw - 0.2
            body = [[(r['title'], dict(bold=True, size=13.5))]]
            if r['desc']: body.append([(r['desc'], dict(size=9, color=MUTED))])
            D.text(s, tx, y + 0.08, tw, ch - 0.16, body, 13, WHITE, anchor='m', spacing=1.12)
            D.text(s, cw + x - prw - 0.2, y, prw, ch, pt, 14, pc, True, align='r', anchor='m')
        by = min(9.6, top + n * (ch + gap) + 0.2)
        D.box(s, 0.5, by, 7.27, 1.0, CARD2, OR, r=0.18, lw=1.0)
        D.text(s, 0.8, by, 6.8, 1.0, [[('Уже в ціні крана:  ', dict(color=MUTED, size=11)),
                                         (f"ротатор {cs['rotator_t']} т  +  грейфер {cs['grapple_m2']} м²", dict(bold=True, size=15))]], 15, WHITE, anchor='m')

    def tile(self, r, w, h):
        iid = r.get('item_id') or ''
        if r.get('kind') == 'trailer' and self.npl:
            return self.raster('tile-npl', lambda p_, w_, h_: jpg(R.on_card(Image.open(R.asset(p_)).convert('RGBA'), w_, h_, 0.95)), NPL['cutout'], w, h, assets=(NPL['cutout'],))
        if r.get('kind') == 'crane':
            return self.raster('tile-crane', lambda c, w_, h_: jpg(R.on_card(Image.open(io.BytesIO(c)), w_, h_, 0.95)), self.cut, w, h)
        path = next((v for k, v in TILES.items() if v and iid.startswith(k)), None) or r.get('photo')
        if not path or not R.asset_ok(path):
            m = re.search(r'(\d+)\s*(л|шт|т)\b', r.get('title') or '')
            return self.raster('tile-none', no_photo, m.group(0) if m else 'BMF', w, h)
        tw = CARD_TWEAKS.get(path, {})
        return self.raster('tile', lambda p, w_, h_, t: jpg(R.card_photo(p, w_, h_, **dict(t))), path, w, h, tuple(tw.items()), assets=(path,))

    # 06 options
    def p_options(self):
        D, sp = self.deck, self.spec
        accs = sp.get('accessories') or []
        if not accs: return
        s = D.slide(self.raster('bg-a', page_bg, ((8.3, 6.5, 3.2, 0.14),)))
        self.pages.append((s, False))
        self.kicker(s, 0.5, 0.95, 'додаткові опції')
        D.rt(s, self.headline([[('Додайте', WHITE)], [('під свою роботу', OR)]]), 0.48, 1.3, 0.42)
        cols = 2 if len(accs) <= 4 else 3; gap = 0.15
        cw = (7.27 - gap * (cols - 1)) / cols; rn = math.ceil(len(accs) / cols)
        ch = min(3.9 if cols == 2 else 2.75, (10.55 - 2.55 - gap * (rn - 1)) / rn)
        for i, a in enumerate(accs[:9]):
            x, y = 0.5 + (i % cols) * (cw + gap), 2.55 + (i // cols) * (ch + gap)
            D.box(s, x, y, cw, ch, CARD, LINE, r=0.18)
            ph = ch * (0.52 if cols == 2 else 0.42)
            if a.get('photo') and R.asset_ok(a['photo']):
                tw = CARD_TWEAKS.get(a['photo'], {})
                data = self.raster('tile', lambda p, w_, h_, t: jpg(R.card_photo(p, w_, h_, **dict(t))), a['photo'], cw - 0.3, ph, tuple(tw.items()), assets=(a['photo'],))
                D.img(s, data, x + 0.15, y + 0.15, cw - 0.3, ph, r=0.12)
            ty = y + 0.15 + ph + 0.15
            body = [[(a['title'], dict(bold=True, size=12.5 if cols == 2 else 10.5, color=WHITE))]]
            if a.get('desc'): body.append([(a['desc'], dict(size=9 if cols == 2 else 8, color=MUTED))])
            D.text(s, x + 0.22, ty, cw - 0.44, ch - ph - 0.75, body, 10, WHITE, spacing=1.12)
            D.text(s, x + 0.22, y + ch - 0.38, cw - 0.44, 0.25, a['price_text'], 9.5, OR, True)

    # 07 price + terms
    def p_price(self):
        D, sp = self.deck, self.spec
        s = D.slide(self.raster('bg-c', page_bg, ((8.4, 2.4, 3.6, 0.2), (-0.4, 10.5, 3.0, 0.1))))
        self.pages.append((s, False))
        self.kicker(s, 0.5, 0.95, 'вартість')
        D.rt(s, self.headline([[('Вартість', WHITE)], [('комплекту', OR)]]), 0.48, 1.3, 0.42)
        rows = sp['rows']; n = len(rows)
        rh = 0.72 if n <= 5 else max(0.5, 3.9 / n)
        top = 2.55; th = n * rh + 0.3
        D.box(s, 0.5, top, 7.27, th + 1.55, CARD, LINE, r=0.2)
        for i, r in enumerate(rows):
            y = top + 0.15 + i * rh
            if i: D.box(s, 0.8, y, 6.67, 0.01, LINE, r=0)
            D.text(s, 0.8, y + 0.12, 0.45, 0.3, f'{i + 1:02d}', 10, OR, True)
            D.text(s, 1.3, y + 0.08, 4.4, 0.3, r['title'], 14, WHITE, True)
            if r['desc'] and rh >= 0.6: D.text(s, 1.3, y + 0.38, 4.4, 0.28, r['desc'], 9.5, MUTED)
            pt = r['price_text'] or 'у комплекті'
            D.text(s, 5.0, y + 0.08, 2.47, 0.3, pt, 14, WHITE if any(c.isdigit() for c in pt) else MUTED, True, 'r')
        y = top + th + 0.05
        D.box(s, 0.8, y, 6.67, 0.015, OR, r=0)
        D.text(s, 0.8, y + 0.2, 3, 0.25, 'РАЗОМ, З ПДВ', 8.5, MUTED, True, track=2.0)
        total = sp['total_text'].replace(' ₴', '').replace(' ₴', '').strip()
        digits = any(c.isdigit() for c in total)
        hl = self.headline([[(total, WHITE), (' ₴' if digits else '', OR)]], F_BLACK)
        D.rt(s, hl, 0.78, y + 0.52, 0.55 if digits else 0.4, maxw=6.7)
        note = [f"Станом на {sp.get('price_date') or sp['date']}."]
        if sp.get('valid_until'): note.append(f"Пропозиція дійсна до {sp['valid_until']}.")
        if sp.get('price_note'): note.append(sp['price_note'])
        D.text(s, 0.8, y + 1.15, 6.67, 0.35, ' '.join(note), 8, DIM, spacing=1.15)
        terms = sp['terms'][:6]
        ty = top + th + 1.85
        if terms and ty < 9.6:
            D.text(s, 0.5, ty, 4, 0.2, 'УМОВИ СПІВПРАЦІ', 7.5, OR, True, track=2.0)
            cols = 3 if len(terms) > 4 else 2; gap = 0.12; cw = (7.27 - gap * (cols - 1)) / cols
            rn = math.ceil(len(terms) / cols); chh = min(1.25, (10.85 - ty - 0.3 - gap * (rn - 1)) / rn)
            for i, t in enumerate(terms):
                x, yy = 0.5 + (i % cols) * (cw + gap), ty + 0.3 + (i // cols) * (chh + gap)
                D.box(s, x, yy, cw, chh, CARD, LINE, r=0.14)
                D.text(s, x + 0.18, yy + 0.14, cw - 0.36, 0.25, t['title'], 10.5, WHITE, True)
                D.text(s, x + 0.18, yy + 0.42, cw - 0.36, chh - 0.5, t['text'], 8.5, MUTED, spacing=1.15)

    # 08 contacts on a full-bleed photo
    def p_end(self):
        D, sp = self.deck, self.spec
        s = D.slide(self.raster('bg-d', page_bg, ((4.1, 11.9, 4.5, 0.22),), False))
        self.pages.append((s, True))
        d = self.photo_of('end', NPL['end'] if self.npl else TRUCK_PH['end'] if self.truck else PHOTOS[self.c['profile']]['end'])
        D.img(s, self.raster('photo-end', photo, d[0], W, 7.4, d[1], d[2], 0.42, assets=(d[0],)), 0, 0, W, 7.4)
        D.rt(s, self.headline([[('Замовте', WHITE)], [('НПЛ-10' if self.npl else f'{self.truck_name} + {self.no or self.m}' if self.truck else self.m, OR)]], F_BLACK), 0.47, 5.35, 0.62)
        lead = ('Відповімо на запитання, підберемо кран і комплектацію НПЛ-10 під вашу роботу та підготуємо договір.' if self.npl else
                f'Відповімо на запитання, оглянемо вашу вантажівку {self.truck_name}, підберемо комплектацію і підготуємо договір.' if self.truck else
                'Відповімо на запитання, підберемо комплектацію під ваш причіп або трактор і підготуємо договір.')
        D.text(s, 0.5, 6.95, 5.6, 0.5, lead, 11, (214, 217, 222), spacing=1.2)
        ct = sp['contacts']
        cards = [('ТЕЛЕФОН', ct['phones'][:2]), ('САЙТ', ct['sites'][:2]), ('EMAIL', [ct['email']] if ct['email'] else []),
                 ('АДРЕСА', ct['address'].split('\n') if ct['address'] else [])]
        cw = (7.27 - 0.15) / 2
        for i, (t, lines) in enumerate(cards):
            x, y = 0.5 + (i % 2) * (cw + 0.15), 7.75 + (i // 2) * 1.05
            D.box(s, x, y, cw, 0.92, CARD, LINE, r=0.16)
            D.text(s, x + 0.25, y + 0.14, cw - 0.4, 0.2, t, 7.5, OR, True, track=2.0)
            D.text(s, x + 0.25, y + 0.36, cw - 0.4, 0.5, lines or ['—'], 11 if t != 'АДРЕСА' else 9.5, WHITE, True, spacing=1.1)
        y = 9.95
        D.text(s, 0.5, y + 0.12, 3.6, 0.3, 'Скануйте — і ми на зв’язку', 11, WHITE, True)
        D.text(s, 0.5, y + 0.42, 3.6, 0.3, 'сайт · TikTok · Instagram', 9, MUTED)
        for i, (q, cap) in enumerate(zip(self.qrs, ('САЙТ', 'TIKTOK', 'INSTAGRAM'))):
            x = 5.0 + i * 0.95
            D.box(s, x, y, 0.82, 0.82, WHITE, r=0.1)
            D.img(s, q, x + 0.06, y + 0.06, 0.7, 0.7)
        D.text(s, 0.5, 10.95 - 0.02, 4, 0.2, sp['maker']['line'].upper(), 6.5, DIM, True, track=1.5)


def main(argv=None):
    from .spec import load
    from .pdf import to_pdf
    from .__main__ import file_name
    ap = argparse.ArgumentParser(prog='kpgen.ultra'); ap.add_argument('--spec', required=True); ap.add_argument('--out', required=True)
    ap.add_argument('--no-pdf', action='store_true'); a = ap.parse_args(argv)
    spec = load(a.spec); os.makedirs(a.out, exist_ok=True)
    bd = os.environ.get('KPGEN_BUILD') or os.path.join(a.out, '.build')
    prs = Ultra(spec, bd).build()
    out = os.path.join(os.path.abspath(a.out), file_name(spec) + '.pptx'); prs.save(out)
    pdf = None if a.no_pdf else to_pdf(out, os.path.abspath(a.out))
    print(json.dumps({'pptx': out, 'pdf': pdf, 'pages': len(prs.slides)}, ensure_ascii=False))


if __name__ == '__main__':
    sys.exit(main())
