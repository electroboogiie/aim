# What the admin can edit per style: the fixed texts of each page (texts the engine never rewrites) and the
# current background of each page. python3 -m kpgen.introspect STYLE  -> JSON {page: {title, texts: [...]}}
import io, json, os, re, sys, tempfile
from PIL import Image
from pptx import Presentation

from .spec import validate
from .styles import STYLES, pages, template_path

HERE = os.path.dirname(os.path.abspath(__file__))
SAMPLES = {'v3': ['bmf720_komplekt', 'bmf750_radiator', 'npl10_bmf720'], 'd6': ['npl10_bmf750']}
SKIP = re.compile(r'^([ST]-\d\d|\d\d|\d\d / \d\d|[\d\s,.]+|//.*|→|●.*|[A-Z]|L|W|H|▭)$')


def _texts(slide):
    return {sh.shape_id: sh.text_frame.text for sh in slide.shapes if sh.has_text_frame and sh.text_frame.text.strip()}


def _build(style, name, build_dir):
    with open(os.path.join(HERE, '..', 'tests', 'specs', name + '.json'), encoding='utf-8') as f: spec = json.load(f)
    spec['style'] = style
    # every text the spec feeds gets a marker, so whatever the engine writes differs from the template
    spec['terms'] = [dict(title=f'Q{i}', text=f'Q{i}') for i in range(5)]
    spec['contacts'] = dict(phones=['Q1', 'Q2'], email='Q', sites=['q.q', 'q.w'], address='Q, Q, Q', socials=['Q', 'W'])
    spec['maker'] = dict(name='Q', line='Q'); spec['price_note'] = 'Q'; spec['cover_text'] = 'Q'
    spec['crane'] = dict(spec['crane'], lead='Q', lead2='Q')
    spec = validate(spec)
    if STYLES[style]['engine'] == 'd6':
        from .d6 import D6 as E
    else:
        from .build import Engine as E
    e = E(spec, build_dir); e.build()
    return e.pages


def static_texts(style):
    tpl = list(Presentation(template_path(style)).slides)
    origin = {k: i for k, _, i in pages(style)}
    seen, changed = {}, set()
    with tempfile.TemporaryDirectory() as d:
        for name in SAMPLES[STYLES[style]['engine']]:
            for key, slide in _build(style, name, d).items():
                base = _texts(tpl[origin[key]]); now = _texts(slide)
                seen.setdefault(key, base)
                for sid, t in base.items():
                    if now.get(sid) != t: changed.add((key, sid))
    out = {}
    for key, title, _ in pages(style):
        if key not in seen: continue
        out[key] = dict(title=title, texts=[dict(id=sid, text=t) for sid, t in sorted(seen[key].items())
                                          if (key, sid) not in changed and not SKIP.match(t.strip())])
    return out


def background(style, key, width=330):
    """JPEG thumbnail of the page's background picture in the style's template."""
    origin = {k: i for k, _, i in pages(style)}
    s = list(Presentation(template_path(style)).slides)[origin[key]]
    for sh in s.shapes:
        if sh.name in ('bg', 'bg_grid'):
            im = Image.open(io.BytesIO(sh.image.blob)).convert('RGB'); im.thumbnail((width, width * 2))
            b = io.BytesIO(); im.save(b, 'JPEG', quality=80); return b.getvalue()
    return None


if __name__ == '__main__':
    print(json.dumps(static_texts(sys.argv[1]), ensure_ascii=False, indent=1))
