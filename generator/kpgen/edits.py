# Page editor support, shared by the KP_2_v3 and Ultra engines.
#   spec.edit = {hide_pages: [page key], shapes: [{page, sid, was, text?|hide?|path?}]}
# A shape is found by its id and checked by `was` (its text, or 'pic:<hash>' for a picture); when another edit moved the
# ids (a position added), the first shape on the page with that text / picture is taken. An edit that finds nothing is skipped.
# zones(): what the editor shows over the page images: positions (row:<item_id>), the «add position» spot, texts, photos.
import hashlib, io
from PIL import Image
from pptx.enum.shapes import MSO_SHAPE_TYPE

from . import raster as R
from .pptx_util import EMU, set_image, remove

PAGE_TITLES = dict(cover='Титулка', conversion='Дообладнання', truck='Лісовоз', manifest='Кран', anatomy='Конструкція',
                   passport='Паспорт крана', trailer='Причіп', trailer_spec='Габарити причепа', npl='НПЛ-10',
                   npl_spec='Характеристики НПЛ-10', crane='Кран', drawings='Креслення', chart='Діаграма вантажу',
                   config='Конфігурація', kit='Склад комплекту', options='Додаткові опції', price='Вартість',
                   maker='Виробник BMF', own_production='Власне виробництво', own='Власне виробництво', end='Контакти')


def hidden_pages(spec):
    return {k for k in ((spec.get('edit') or {}).get('hide_pages') or []) if isinstance(k, str) and k != 'cover'}


def drop_slide(prs, slide):
    ids = prs.slides._sldIdLst
    for el in list(ids):
        if int(el.get('id')) == slide.slide_id:
            prs.part.drop_rel(el.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id'))
            ids.remove(el)
            return


def pic_hash(sh):
    try: return 'pic:' + hashlib.sha1(sh.image.blob).hexdigest()[:10]
    except Exception: return None  # noqa: BLE001 - linked or empty picture


def is_pic(sh):
    return sh.shape_type == MSO_SHAPE_TYPE.PICTURE or (getattr(sh, '_element', None) is not None and sh._element.tag.endswith('}pic'))


def was_of(sh):
    if sh.name.startswith('hl:'): return sh.name   # Ultra headline picture: its original text
    if is_pic(sh): return pic_hash(sh)
    if sh.has_text_frame: return sh.text_frame.text
    return None


def set_text(sh, text):
    """Replace the shape's text, one paragraph per line, each keeping the formatting of the first run of its paragraph."""
    import copy
    tf = sh.text_frame; ps = list(tf.paragraphs); lines = str(text).split('\n') or ['']
    while len(ps) < len(lines):
        el = copy.deepcopy(ps[-1]._p); ps[-1]._p.addnext(el); ps = list(tf.paragraphs)
    for p in ps[len(lines):]: tf._txBody.remove(p._p)
    for p, line in zip(tf.paragraphs, lines):
        runs = p.runs
        if not runs: p.add_run().text = line; continue
        runs[0].text = line
        for r in runs[1:]: r._r.getparent().remove(r._r)


def photo_into(sh, path, fx=0.5, fy=0.5):
    """The asset cropped to the picture's box (aspect of the box, focus fx, fy)."""
    w, h = sh.width / EMU, sh.height / EMU
    im = Image.open(R.asset(path)); alpha = im.mode in ('RGBA', 'LA', 'P')
    im = im.convert('RGBA' if alpha else 'RGB')
    r = max(w / im.width, h / im.height); cw, ch = round(w / r), round(h / r)
    x, y = int((im.width - cw) * fx), int((im.height - ch) * fy)
    im = im.crop((x, y, x + cw, y + ch))
    tw = min(im.width, int(w * 200)); im = im.resize((tw, max(1, int(tw * h / w))), Image.LANCZOS)
    b = io.BytesIO(); im.save(b, 'PNG' if alpha else 'JPEG', **({} if alpha else {'quality': 88})); set_image(sh, b.getvalue())


def apply(pages, spec):
    """pages: {key: slide}. Returns {(key, sid): original was} of the shapes it changed (the editor keeps editing them by it)."""
    done = {}
    for e in (spec.get('edit') or {}).get('shapes') or []:
        if not e.get('was'): continue
        if e.get('page') == '*':   # the same text on every page that has it (the KP title in the page headers)
            targets = [(k, [x for x in s.shapes if was_of(x) == e['was']]) for k, s in pages.items()]
        else:
            s = pages.get(e.get('page'))
            if s is None: continue
            sh = next((x for x in s.shapes if x.shape_id == e.get('sid') and was_of(x) == e['was']), None) \
                or next((x for x in s.shapes if was_of(x) == e['was']), None)
            targets = [(e['page'], [sh] if sh is not None else [])]
        for key, shapes in targets:
            for sh in shapes:
                if e.get('hide'): remove(sh); continue
                if sh.name.startswith('hl:'): pass   # headline picture: drawn with the new text by the engine
                elif e.get('text') is not None and sh.has_text_frame: set_text(sh, e['text'])
                elif e.get('path') and is_pic(sh) and R.asset_ok(e['path']): photo_into(sh, e['path'], float(e.get('fx', 0.5)), float(e.get('fy', 0.5)))
                else: continue
                done[(key, sh.shape_id)] = e['was']
    return done


def _box(sh):
    return [sh.left / EMU, sh.top / EMU, sh.width / EMU, sh.height / EMU]


def _font_pt(sh):
    best = 0
    for p in sh.text_frame.paragraphs:
        for r in p.runs:
            if r.font.size: best = max(best, r.font.size.pt)
    return best or 10


def zones(prs, order, reg, done, style, page_w, page_h, hero=None, dropped=()):
    """order: [(key, slide)] in page order; reg: [dict(s=slide, id, kind, what=[shapes] or (x, y, w, h) inches, item?)]."""
    out = []
    for key, s in order:
        zs = []
        for r_ in reg:
            if r_['s'] is not s: continue
            what = r_['what']
            if isinstance(what, tuple): b = list(what)
            else:
                live = [x for x in what if x is not None and x._element.getparent() is not None]
                if not live: continue
                bs = [_box(x) for x in live]
                x0, y0 = min(b[0] for b in bs), min(b[1] for b in bs)
                b = [x0, y0, max(b[0] + b[2] for b in bs) - x0, max(b[1] + b[3] for b in bs) - y0]
            zs.append(dict({k: v for k, v in r_.items() if k not in ('s', 'what')}, box=b))
        rows = [z for z in zs if z['kind'] in ('row', 'tile')]
        for sh in s.shapes:
            b = _box(sh)
            if b[2] <= 0.05 or b[3] <= 0.02: continue
            if sh.name.startswith('hl:'):   # Ultra headline drawn as a picture: edited as a text
                text = sh._element.nvPicPr.cNvPr.get('descr') or sh.name[3:]
                z = dict(id=f's{sh.shape_id}', kind='text', sid=sh.shape_id, box=b, text=text, was=sh.name, pt=30, edited=text != sh.name[3:])
            elif is_pic(sh):
                if b[2] < 0.9 or b[3] < 0.6: continue   # icons, logos, QR codes
                z = dict(id=f's{sh.shape_id}', kind='pic', sid=sh.shape_id, box=b, was=done.get((key, sh.shape_id)) or pic_hash(sh),
                         full=b[2] > page_w * 0.9 and b[3] > page_h * 0.5)
                hk = hero(key, sh) if hero else None
                if hk: z['hero'] = hk   # page photo: changed through spec.overrides.bg (cropped and toned as the style does)
            elif sh.has_text_frame and sh.text_frame.text.strip():
                t = sh.text_frame.text
                z = dict(id=f's{sh.shape_id}', kind='text', sid=sh.shape_id, box=b, text=t, was=done.get((key, sh.shape_id)) or t,
                         pt=_font_pt(sh), edited=(key, sh.shape_id) in done)
            else: continue
            cx, cy = b[0] + b[2] / 2, b[1] + b[3] / 2
            row = None if z.get('full') else next((r['id'] for r in rows if r['box'][0] <= cx <= r['box'][0] + r['box'][2] and r['box'][1] <= cy <= r['box'][1] + r['box'][3]), None)
            if row: z['row'] = row
            zs.append(z)
        for z in zs: z['box'] = [round(z['box'][0] / page_w, 4), round(z['box'][1] / page_h, 4), round(z['box'][2] / page_w, 4), round(z['box'][3] / page_h, 4)]
        out.append(dict(key=key, title=PAGE_TITLES.get(key, key), zones=zs))
    return dict(style=style, pages=out, hidden=[dict(key=k, title=PAGE_TITLES.get(k, k)) for k in dropped])
