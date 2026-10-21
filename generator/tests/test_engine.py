# Engine checks without PDF export:  cd generator && python3 -m unittest discover -s tests -v
import copy, json, os, re, shutil, subprocess, sys, tempfile, unittest
from pptx import Presentation

HERE = os.path.dirname(os.path.abspath(__file__))
GEN = os.path.dirname(HERE)
sys.path.insert(0, HERE)
import build_samples as B  # noqa: E402

ENV = dict(os.environ, KPGEN_BUILD=os.path.join(HERE, 'out', '.build'))


TMP = tempfile.mkdtemp(prefix='kpgen_test_')


def tearDownModule(): shutil.rmtree(TMP, ignore_errors=True)


def run(spec):
    d = tempfile.mkdtemp(dir=TMP)
    sp = os.path.join(d, 'spec.json')
    with open(sp, 'w', encoding='utf-8') as f: json.dump(spec, f, ensure_ascii=False)
    p = subprocess.run([sys.executable, '-m', 'kpgen', '--spec', sp, '--out', d, '--no-pdf'], cwd=GEN, env=ENV,
                       stdout=subprocess.PIPE, stderr=subprocess.PIPE, text=True)
    res = json.loads(p.stdout.strip().splitlines()[-1]) if p.returncode == 0 else None
    return p.returncode, res, p.stderr


def texts(prs):
    out = []
    for s in prs.slides:
        out.append([sh.text_frame.text for sh in s.shapes if sh.has_text_frame])
    return out


class EngineTest(unittest.TestCase):
    def sample(self, name): return copy.deepcopy(B.SAMPLES[name][0])

    def check_numbers(self, prs):
        T = texts(prs); n = len(T)
        for i, page in enumerate(T, 1):
            nums = [t for t in page if re.fullmatch(r'\d\d / \d\d', t)]
            if i > 1: self.assertEqual(nums, [f'{i:02d} / {n:02d}'], f'page {i}')
            tags = [t for t in page if re.match(r'// \d\d — ', t)]
            for t in tags: self.assertTrue(t.startswith(f'// {i - 1:02d} — '), f'page {i}: {t}')
        return T

    def test_npl_without_accessories(self):
        spec = self.sample('npl10_bmf720'); spec['accessories'] = []
        code, res, err = run(spec)
        self.assertEqual(code, 0, err); self.assertEqual(res['pages'], 10)
        T = self.check_numbers(Presentation(res['pptx']))
        config = T[6]
        self.assertFalse(any('Опції для крана' in t for t in config))
        self.assertFalse(any('ДОДАТКОВІ ОПЦІЇ' in t for p in T for t in p))

    def test_npl_rows_have_no_hydraulics(self):
        code, res, err = run(self.sample('npl10_bmf750'))
        self.assertEqual(code, 0, err); self.assertEqual(res['pages'], 11)
        T = self.check_numbers(Presentation(res['pptx']))
        for page in (T[6], T[8]):   # configuration and price pages
            self.assertFalse(any('автономна гідравліка' in t.lower() for t in page))
            self.assertEqual(sum(t == 'у комплекті' for t in page), 2)
        self.assertTrue(any('1 978 000 ₴' in t for t in T[8]))
        self.assertTrue(any('ДІЙСНА ДО 25.10.2026' in t for t in T[8]))

    def test_crane_kit_variants(self):
        for n in (1, 3, 5, 8):
            spec = self.sample('bmf750_radiator')
            rows = (spec['rows'] * 2)[:n]
            for i, r in enumerate(rows, 1): r['no'] = i
            spec['rows'] = rows
            code, res, err = run(spec)
            self.assertEqual(code, 0, f'{n} rows: {err}')
            T = self.check_numbers(Presentation(res['pptx']))
            price = T[6]
            self.assertEqual(sum(1 for t in price if re.fullmatch(r'0\d', t)), n, f'{n} rows')

    def test_heavy_label_crane_kit(self):
        """«Heavy Label» for a crane in a kit: 8 pages renumbered, the model everywhere, its chip orange, no НПЛ-10 kit rows."""
        for name, model in (('bmf650_komplekt', 'BMF 650'), ('bmf750_radiator', 'BMF 750'), ('bmf720_komplekt', 'BMF 720')):
            spec = self.sample(name); spec['style'] = 'crane_label'
            code, res, err = run(spec)
            self.assertEqual(code, 0, f'{name}: {err}'); self.assertEqual(res['pages'], 8)
            prs = Presentation(res['pptx']); T = texts(prs)
            for i, page in enumerate(T[1:], 2):
                self.assertIn(f'{i:02d} / 08', page, f'{name} page {i}')
                self.assertIn(f'{model} · КРАН У КОМПЛЕКТІ', page)
            self.assertFalse(any('BMF 720' in t for p in T for t in p if model != 'BMF 720' and not t.startswith('BMF 720\n')
                                 and t != 'BMF 720'), name)
            chips = {sh.text_frame.text: 'F2530A' in prs.slides[1].shapes._spTree.xpath(f'.//p:cNvPr[@id="{sh.shape_id - 1}"]/../..')[0].xml
                     for sh in prs.slides[1].shapes if sh.shape_id in (46, 48, 50, 52)}
            self.assertEqual([k for k, v in chips.items() if v], [model], name)
            self.assertEqual(set(chips), {'BMF 650', 'BMF 720', 'BMF 750', 'BMF 905'})
            self.assertFalse(any('Напівпричіп' in t for t in T[4]), name)   # price page
        spec = self.sample('npl10_bmf720'); spec['style'] = 'crane_label'
        code, _, err = run(spec)
        self.assertNotEqual(code, 0); self.assertIn('кран у комплекті', err)

    def test_many_accessories_and_terms(self):
        spec = self.sample('npl10_bmf650')
        spec['accessories'] = spec['accessories'] + spec['accessories'][:3]
        spec['terms'] = spec['terms'][:3]
        code, res, err = run(spec)
        self.assertEqual(code, 0, err)
        T = self.check_numbers(Presentation(res['pptx']))
        opts = T[7]
        self.assertEqual(sum('ціна за запитом' == t for t in opts), 9)
        self.assertTrue(any('Також на замовлення: Регулюючий клапан, Грейфери Hardox, Ротатори.' in t for t in opts))
        self.assertEqual(sum(bool(re.fullmatch(r'T-0\d', t)) for t in T[8]), 3)

    def test_bad_photo_paths(self):
        spec = self.sample('bmf650plus_2pos')
        spec['photos'] = dict(cover='../kpgen/__init__.py', end='bmf/немає.jpg')
        spec['rows'][1]['photo'] = '/etc/hostname'
        spec['accessories'] = [dict(group='crane', title='Ротатори', desc='', photo='../../README.md', price_text='')]
        code, res, err = run(spec)
        self.assertEqual(code, 0, err); self.assertEqual(res['pages'], 9)
        self.assertEqual(err.count('увага'), 4, err)
        spec['photos'] = dict(cover='bmf/650_a.jpg', cover_focus=[0.5, 2])
        code, _, err = run(spec); self.assertEqual(code, 1); self.assertIn('cover_focus', err)

    def test_errors(self):
        spec = self.sample('bmf720_komplekt'); spec['crane']['profile'] = 'bmf850'
        code, _, err = run(spec); self.assertEqual(code, 2); self.assertIn('профілю', err)
        spec = self.sample('bmf720_komplekt'); spec['crane'] = None
        code, _, err = run(spec); self.assertEqual(code, 2)
        spec = self.sample('bmf720_komplekt'); del spec['rows']
        code, _, err = run(spec); self.assertEqual(code, 1); self.assertIn('rows', err)
        spec = self.sample('bmf720_komplekt'); del spec['crane']['specs']['reach_m']
        code, _, err = run(spec); self.assertEqual(code, 1); self.assertIn('reach_m', err)
        spec = self.sample('bmf720_komplekt'); spec['rows'] = spec['rows'] * 3
        code, _, err = run(spec); self.assertEqual(code, 1); self.assertIn('максимум 8', err)


if __name__ == '__main__':
    unittest.main()
